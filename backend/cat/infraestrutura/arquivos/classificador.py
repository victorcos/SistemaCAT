"""Descobrir o que cada arquivo de uma pasta é, sem ler o arquivo inteiro.

A pasta de trabalho de uma empresa é bagunçada por natureza: EFD ICMS/IPI ao
lado de EFD Contribuições e de ECD, XML de NF-e junto com XML de evento,
relatório do ERP, planilha de conferência, zip que ninguém abriu, e arquivo de
outro cliente que caiu ali. Perguntar ao usuário o que é cada coisa não escala
com 7.036 arquivos, e confiar na extensão ou no nome não funciona — todos se
chamam `.txt`.

Então o tipo sai do **conteúdo**, e sai barato: o SPED se identifica na
primeira linha, o XML nos primeiros bytes, e o relatório gerencial no
cabeçalho. Nenhum arquivo é lido inteiro nesta etapa.
"""

from __future__ import annotations

import hashlib
import os
import re
import unicodedata
import zipfile
from datetime import date
from typing import Iterator

from cat.dominio.lote import (
    ArquivoDoLote,
    abas_de_canceladas,
    CertificadosIgnorados,
    TipoDeArquivo,
    e_arquivo_de_certificado,
    e_pasta_de_certificado,
)
from cat.dominio.gerencial.campos import Especie
from cat.dominio.sped.cabecalho import (
    ArquivoNaoReconhecido,
    CabecalhoSped,
    TipoSped,
    ler_cabecalho,
)
from cat.infraestrutura.arquivos.gerencial import inspecionar
from cat.infraestrutura.arquivos.xml_compactado import e_zip, membros_xml
from cat.log import obter_log

log = obter_log(__name__)

CODIFICACOES = ("utf-8-sig", "utf-8", "latin-1")
# 16 KB e não 8: no XML de NF-e o bloco <dest> só vem depois do <emit> com o
# endereço inteiro, e 8 KB nem sempre alcançava — a nota ficava sem
# destinatário e uma compra podia passar por "de outra empresa". Para SPED e
# relatório o extra não muda nada: a primeira linha cabe de sobra.
BYTES_DE_AMOSTRA = 16384

EXTENSOES_TEXTO = (".txt", ".sped", ".efd")
EXTENSOES_XML = (".xml",)
EXTENSOES_COMPACTADAS = (".zip", ".rar", ".7z", ".gz", ".tar", ".tgz")

# uma varredura tem de terminar: pasta de rede com milhões de XML existe
LIMITE_DE_ARQUIVOS = 200_000

_POR_TIPO_SPED = {
    TipoSped.EFD_ICMS_IPI: TipoDeArquivo.SPED_ICMS_IPI,
    TipoSped.EFD_CONTRIBUICOES: TipoDeArquivo.SPED_CONTRIBUICOES,
    TipoSped.ECD: TipoDeArquivo.SPED_ECD,
    TipoSped.ECF: TipoDeArquivo.SPED_ECF,
}

_POR_ESPECIE = {
    Especie.MOVIMENTO: TipoDeArquivo.GERENCIAL_MOVIMENTO,
    Especie.INVENTARIO: TipoDeArquivo.GERENCIAL_INVENTARIO,
    Especie.RESUMO: TipoDeArquivo.GERENCIAL_RESUMO,
}

# Marca de erro de sincronização do OneDrive: o arquivo de verdade nunca desceu
# e no lugar dele fica um texto de 135 bytes. Numa pasta real desta casa, 149
# de 284 arquivos eram isto — quem contasse arquivo acharia que tinha a base
# inteira. Precisa ter nome e tamanho e conteúdo: nome sozinho pega arquivo
# legítimo que por acaso termine assim.
TAMANHO_MAXIMO_DE_STUB = 2048
_RE_STUB = re.compile(
    r"exceeds the allowed limit|CorrelationId", re.IGNORECASE)

# o XML de NF-e, NFC-e ou CF-e SAT: a chave e o emitente aparecem nos primeiros
# KB. O cupom SAT entra desde a v0.53, quando a etapa 3 passou a ler o item do XML
_RE_NFE = re.compile(r"<(nfeProc|NFe|infNFe|CFe|infCFe)\b", re.IGNORECASE)
_RE_CNPJ_EMITENTE = re.compile(
    r"<emit>.*?<CNPJ>(\d{14})</CNPJ>", re.IGNORECASE | re.DOTALL)
# Numa nota que a empresa RECEBE, o emitente é o fornecedor. Sem o
# destinatário, todo XML de compra — o insumo principal da CAT 42 — seria
# jogado fora como "de outra empresa". O bloco <dest> pode trazer CPF em
# vez de CNPJ (venda a consumidor); aí não há o que capturar, e tudo bem.
_RE_CNPJ_DESTINATARIO = re.compile(
    r"<dest>.*?<CNPJ>(\d{14})</CNPJ>", re.IGNORECASE | re.DOTALL)
# aaaa-mm-dd na NF-e, aaaammdd no CF-e
_RE_EMISSAO = re.compile(r"<(?:dhEmi|dEmi)>(\d{4})-?(\d{2})-?(\d{2})")
# evento de cancelamento de NF-e: tira a nota da movimentação na etapa 3
_RE_EVENTO_CANCELAMENTO = re.compile(r"<tpEvento>\s*110111\s*</tpEvento>")
# o cancelamento de CT-e também é 110111; o de NF-e é o que traz <chNFe>
_RE_CHAVE_DO_EVENTO = re.compile(r"<chNFe>\s*(\d{44})\s*</chNFe>")
_RE_CHAVE_SOLTA = re.compile(r"(?<!\d)\d{44}(?!\d)")
# no zip, a nota se reconhece pelo Id: CT-e também tem <infNFe>, dentro do infDoc
_RE_ID_DE_NOTA = re.compile(r'Id\s*=\s*["\'](?:NFe|CFe)\d{44}["\']')
MEMBROS_NA_AMOSTRA_DO_ZIP = 50


def _amostra(caminho: str) -> str:
    with open(caminho, "rb") as f:
        bruto = f.read(BYTES_DE_AMOSTRA)
    for cod in CODIFICACOES:
        try:
            return bruto.decode(cod)
        except UnicodeDecodeError:
            continue
    return bruto.decode("latin-1", errors="replace")


def _do_sped(caminho: str, tamanho: int, texto: str) -> ArquivoDoLote | None:
    linha = next((l for l in texto.splitlines() if l.strip()), "")
    if not linha.startswith("|0000|"):
        return None
    try:
        c: CabecalhoSped = ler_cabecalho(linha)
    except ArquivoNaoReconhecido as erro:
        return ArquivoDoLote(
            caminho=caminho, nome=os.path.basename(caminho), tamanho=tamanho,
            tipo=TipoDeArquivo.SPED_OUTRO, motivo=str(erro),
        )
    return ArquivoDoLote(
        caminho=caminho,
        nome=os.path.basename(caminho),
        tamanho=tamanho,
        tipo=_POR_TIPO_SPED.get(c.tipo, TipoDeArquivo.SPED_OUTRO),
        cnpj=c.cnpj.valor if c.cnpj else None,
        competencia=c.inicio,
        uf=c.uf,
        detalhe=c.nome,
        retificadora=c.retificadora,
    )


def _do_arquivo_da_cat42(caminho: str, tamanho: int, texto: str) -> ArquivoDoLote | None:
    """O arquivo digital da CAT 42: `0000|mmaaaa|NOME|CNPJ|IE|COD_MUN|01|00`.

    Sem `|` no começo — é o que o separa da EFD, que começa com `|0000|`.
    """
    linha = next((l for l in texto.splitlines() if l.strip()), "")
    campos = linha.split("|")
    if len(campos) != 8 or campos[0] != "0000" or len(campos[1]) != 6 or not campos[1].isdigit():
        return None
    mes, ano = int(campos[1][:2]), int(campos[1][2:])
    if not 1 <= mes <= 12:
        return None
    return ArquivoDoLote(
        caminho=caminho, nome=os.path.basename(caminho), tamanho=tamanho,
        tipo=TipoDeArquivo.CAT42_ARQUIVO_DIGITAL, cnpj=campos[3] or None,
        competencia=date(ano, mes, 1), uf="SP", detalhe=campos[2],
    )


def _competencia_da_chave(chave: str) -> date | None:
    """AAMM das posições 3 a 6 da chave de acesso."""
    ano, mes = int(chave[2:4]), int(chave[4:6])
    return date(2000 + ano, mes, 1) if 1 <= mes <= 12 else None


def _do_xml(caminho: str, tamanho: int, texto: str) -> ArquivoDoLote:
    chave = _RE_CHAVE_DO_EVENTO.search(texto)
    if _RE_EVENTO_CANCELAMENTO.search(texto) and chave and not _RE_NFE.search(texto):
        return ArquivoDoLote(
            caminho=caminho, nome=os.path.basename(caminho), tamanho=tamanho,
            tipo=TipoDeArquivo.XML_CANCELAMENTO,
            cnpj=chave.group(1)[6:20],
            competencia=_competencia_da_chave(chave.group(1)),
            detalhe=f"cancela {chave.group(1)}",
        )
    e_nfe = _RE_NFE.search(texto) is not None
    emitente = _RE_CNPJ_EMITENTE.search(texto)
    destinatario = _RE_CNPJ_DESTINATARIO.search(texto)
    emissao = _RE_EMISSAO.search(texto)
    return ArquivoDoLote(
        caminho=caminho,
        nome=os.path.basename(caminho),
        tamanho=tamanho,
        tipo=TipoDeArquivo.XML_NFE if e_nfe else TipoDeArquivo.XML_OUTRO,
        cnpj=emitente.group(1) if emitente else None,
        cnpj_destinatario=destinatario.group(1) if destinatario else None,
        # a competência de uma nota é o mês da emissão
        competencia=(date(int(emissao.group(1)), int(emissao.group(2)), 1)
                     if emissao else None),
        motivo="" if e_nfe else "o XML não é de NF-e, NFC-e nem CF-e",
    )


def _do_zip(caminho: str, tamanho: int) -> ArquivoDoLote:
    """Zip com XML de nota vira `xml_compactado`; o resto continua `compactado`.

    Lê os diretórios centrais (do zip e dos zips de dentro) e o começo de até 50
    XML. O CNPJ é o da primeira nota achada: é o que separa zip de outra empresa,
    como no XML solto. A competência fica vazia — um zip costuma ter meses inteiros.
    """
    nome = os.path.basename(caminho)
    recusados: list[str] = []
    xmls = 0
    nota = cancelamento = None
    try:
        with zipfile.ZipFile(caminho) as arquivo:
            for _, dono, info in membros_xml(arquivo, caminho, recusados):
                xmls += 1
                if nota is None and xmls <= MEMBROS_NA_AMOSTRA_DO_ZIP:
                    with dono.open(info) as f:
                        texto = f.read(BYTES_DE_AMOSTRA).decode("utf-8", errors="replace")
                    if _RE_ID_DE_NOTA.search(texto):
                        nota = texto
                    elif cancelamento is None and _RE_EVENTO_CANCELAMENTO.search(texto):
                        cancelamento = _RE_CHAVE_DO_EVENTO.search(texto)
    except (OSError, zipfile.BadZipFile, RuntimeError, EOFError, ValueError) as erro:
        log.warning("zip ilegível na classificação", extra={"arquivo": nome, "motivo": str(erro)})
        return ArquivoDoLote(caminho=caminho, nome=nome, tamanho=tamanho,
                             tipo=TipoDeArquivo.COMPACTADO, motivo=f"zip ilegível: {erro}")
    if nota is None and cancelamento is None:
        return ArquivoDoLote(
            caminho=caminho, nome=nome, tamanho=tamanho, tipo=TipoDeArquivo.COMPACTADO,
            motivo=(f"nenhum dos primeiros {min(xmls, MEMBROS_NA_AMOSTRA_DO_ZIP)} XML do zip "
                    "é NF-e, NFC-e ou CF-e") if xmls else "")
    # sem nota na amostra, o CNPJ é o do emitente da nota cancelada, que está na chave
    emitente = _RE_CNPJ_EMITENTE.search(nota) if nota else None
    destinatario = _RE_CNPJ_DESTINATARIO.search(nota) if nota else None
    return ArquivoDoLote(
        caminho=caminho, nome=nome, tamanho=tamanho, tipo=TipoDeArquivo.XML_COMPACTADO,
        cnpj=(emitente.group(1) if emitente else None) if nota else cancelamento.group(1)[6:20],
        cnpj_destinatario=destinatario.group(1) if destinatario else None,
        detalhe=f"{xmls} XML" + (f", {len(recusados)} zip(s) de dentro ilegíveis" if recusados else ""),
    )


def _nome_de_canceladas(nome: str) -> bool:
    sem_acento = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode().lower()
    return "cancel" in sem_acento


def _da_lista_de_canceladas(caminho: str, tamanho: int, texto: str) -> ArquivoDoLote | None:
    """TXT ou CSV com "cancel" no nome e chaves de acesso em pelo menos metade das linhas."""
    if not _nome_de_canceladas(os.path.basename(caminho)):
        return None
    linhas = [l for l in texto.splitlines() if l.strip()]
    com_chave = sum(1 for l in linhas if _RE_CHAVE_SOLTA.search(l))
    if not linhas or com_chave == 0 or com_chave * 2 < len(linhas):
        return None
    return ArquivoDoLote(caminho=caminho, nome=os.path.basename(caminho), tamanho=tamanho,
                         tipo=TipoDeArquivo.LISTA_DE_CANCELADAS,
                         detalhe=f"{com_chave} chaves nas primeiras linhas")


def _da_planilha_de_canceladas(caminho: str, tamanho: int) -> ArquivoDoLote | None:
    """Planilha com "cancel" no nome e alguma chave de acesso escrita como texto."""
    if not _nome_de_canceladas(os.path.basename(caminho)):
        return None
    from openpyxl import load_workbook  # noqa: PLC0415
    try:
        livro = load_workbook(caminho, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as erro:
        log.warning("planilha de canceladas ilegível na classificação",
                    extra={"arquivo": os.path.basename(caminho), "motivo": str(erro)})
        return None
    try:
        chaves = 0
        validas = set(abas_de_canceladas(livro.sheetnames))
        for aba in (a for a in livro.worksheets if a.title in validas):
            for linha in aba.iter_rows(values_only=True, max_row=200):
                chaves += sum(1 for v in linha if isinstance(v, str) and _RE_CHAVE_SOLTA.search(v))
    finally:
        livro.close()
    if not chaves:
        return None
    return ArquivoDoLote(caminho=caminho, nome=os.path.basename(caminho), tamanho=tamanho,
                         tipo=TipoDeArquivo.LISTA_DE_CANCELADAS,
                         detalhe=f"{chaves} chaves nas primeiras linhas")


def _do_gerencial(caminho: str, tamanho: int) -> ArquivoDoLote:
    try:
        insp = inspecionar(caminho)
    except (OSError, ValueError) as erro:
        return ArquivoDoLote(
            caminho=caminho, nome=os.path.basename(caminho), tamanho=tamanho,
            tipo=TipoDeArquivo.DESCONHECIDO, motivo=str(erro),
        )
    if not insp.utilizavel:
        faltam = ", ".join(c.nome for c in insp.mapeamento.faltam_obrigatorios)
        return ArquivoDoLote(
            caminho=caminho, nome=os.path.basename(caminho), tamanho=tamanho,
            tipo=TipoDeArquivo.DESCONHECIDO,
            motivo=f"nem SPED nem relatório: faltam {faltam}",
        )
    return ArquivoDoLote(
        caminho=caminho,
        nome=os.path.basename(caminho),
        tamanho=tamanho,
        tipo=_POR_ESPECIE[insp.especie],
        detalhe=(f"{insp.leiaute.campos} colunas"
                 + (", traz ST do XML" if insp.mapeamento.tem_valores_do_xml
                    else "")),
    )


def classificar(caminho: str, tamanho: int | None = None) -> ArquivoDoLote:
    """O que este arquivo é. Nunca levanta: arquivo ilegível vira desconhecido.

    `tamanho` vem da listagem do diretório quando quem chama já o tem. Em disco
    de rede isso economiza uma ida e volta por arquivo, e numa pasta de 7.036
    SPED a conta é de milhares delas.
    """
    nome = os.path.basename(caminho)
    minusculo = nome.lower()
    if tamanho is None:
        try:
            tamanho = os.path.getsize(caminho)
        except OSError as erro:
            return ArquivoDoLote(caminho=caminho, nome=nome, tamanho=0,
                                 tipo=TipoDeArquivo.DESCONHECIDO, motivo=str(erro))

    if e_zip(caminho):
        return _do_zip(caminho, tamanho)
    if minusculo.endswith(EXTENSOES_COMPACTADAS):
        return ArquivoDoLote(caminho=caminho, nome=nome, tamanho=tamanho,
                             tipo=TipoDeArquivo.COMPACTADO)

    if minusculo.endswith((".xlsx", ".xlsm")):
        planilha = _da_planilha_de_canceladas(caminho, tamanho)
        if planilha is not None:
            return planilha

    try:
        texto = _amostra(caminho)
    except OSError as erro:
        return ArquivoDoLote(caminho=caminho, nome=nome, tamanho=tamanho,
                             tipo=TipoDeArquivo.DESCONHECIDO, motivo=str(erro))

    if (tamanho <= TAMANHO_MAXIMO_DE_STUB
            and minusculo.endswith("_error.txt")
            and _RE_STUB.search(texto)):
        return ArquivoDoLote(caminho=caminho, nome=nome, tamanho=tamanho,
                             tipo=TipoDeArquivo.NAO_BAIXADO,
                             motivo=texto.strip().splitlines()[0][:120])

    if minusculo.endswith(EXTENSOES_XML) or texto.lstrip().startswith("<?xml"):
        return _do_xml(caminho, tamanho, texto)

    do_sped = _do_sped(caminho, tamanho, texto)
    if do_sped is not None:
        return do_sped

    da_cat42 = _do_arquivo_da_cat42(caminho, tamanho, texto)
    if da_cat42 is not None:
        return da_cat42

    if minusculo.endswith(EXTENSOES_TEXTO) or minusculo.endswith(".csv"):
        lista = _da_lista_de_canceladas(caminho, tamanho, texto)
        if lista is not None:
            return lista
        return _do_gerencial(caminho, tamanho)

    return ArquivoDoLote(caminho=caminho, nome=nome, tamanho=tamanho,
                         tipo=TipoDeArquivo.DESCONHECIDO,
                         motivo="extensão não tratada")


def hash_de(caminho: str, bloco: int = 1 << 20) -> str:
    """SHA-256 do arquivo inteiro, lido em blocos de 1 MB.

    É a única leitura completa que a importação faz, e por isso só é chamada
    para candidatos a cópia. Um SPED de 14 MB leva bem menos de um segundo em
    disco local; em rede é o tempo de transferir o arquivo.
    """
    resumo = hashlib.sha256()
    with open(caminho, "rb") as f:
        for pedaco in iter(lambda: f.read(bloco), b""):
            resumo.update(pedaco)
    return resumo.hexdigest()


def percorrer_pasta(pasta: str,
                    certificados: CertificadosIgnorados | None = None) -> Iterator[tuple[str, int]]:
    """(caminho, tamanho) de cada arquivo, entrando nas subpastas.

    Usa `scandir` em vez de `os.walk` porque a listagem do diretório já traz o
    tamanho: pegá-lo aqui evita um `getsize` por arquivo, que em disco de rede
    é uma viagem inteira até o servidor.

    Pasta de certificado não é aberta e .pfx/.p12 não sai daqui: só a contagem
    vai para `certificados`, e o log não leva nome.
    """
    certificados = certificados if certificados is not None else CertificadosIgnorados()
    vistos = 0
    pilha = [pasta]
    while pilha:
        atual = pilha.pop()
        try:
            with os.scandir(atual) as entradas:
                for entrada in sorted(entradas, key=lambda e: e.name):
                    try:
                        if entrada.is_dir(follow_symlinks=False):
                            if e_pasta_de_certificado(entrada.name):
                                certificados.pastas += 1
                            else:
                                pilha.append(entrada.path)
                            continue
                        if e_arquivo_de_certificado(entrada.name):
                            certificados.arquivos += 1
                            continue
                        tamanho = entrada.stat(follow_symlinks=False).st_size
                    except OSError:
                        tamanho = 0
                    vistos += 1
                    if vistos > LIMITE_DE_ARQUIVOS:
                        log.warning(
                            "varredura interrompida no limite",
                            extra={"pasta": pasta, "limite": LIMITE_DE_ARQUIVOS},
                        )
                        return
                    yield entrada.path, tamanho
        except OSError as erro:
            # subpasta sem permissão não derruba a varredura inteira
            log.warning("subpasta ignorada",
                        extra={"subpasta": atual, "motivo": str(erro)})
    if certificados.pastas or certificados.arquivos:
        log.info("certificados deixados de fora da varredura, sem abrir",
                 extra={"pasta": pasta, "pastas": certificados.pastas, "arquivos": certificados.arquivos})
