"""As duas planilhas que alimentam a simulação da reforma.

**A base custeada** — as saídas do mês com o custo de compra de cada linha, no
leiaute "Margem e Impacto Tributário" da primeira simulação (empresa 23): loja,
item, receita bruta, débitos de hoje, custo líquido, IPI e PIS/COFINS sem
crédito, créditos de ICMS e PIS/COFINS. Uma linha por loja × item × CFOP. O
custeio — achar o custo médio de compra de cada venda — ainda é feito fora do
sistema; quando entrar, ele produz esta mesma base.

**A classificação dos itens** — a revisão de SKU com NCM revisada, CST e
cClassTrib, opcionalmente com a categoria do Imposto Seletivo marcada à mão e a
categoria mercadológica do cliente.

As duas se reconhecem pelo cabeçalho, e não pelo nome: o nome do arquivo é do
cliente e muda a cada remessa. O cabeçalho da classificação não está na
primeira linha (vem depois do título, na linha 10 da primeira remessa), por
isso a busca percorre as primeiras linhas.
"""

from __future__ import annotations

import os
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Callable, Iterator

from cat.log import obter_log

log = obter_log(__name__)

# o título da classificação ocupa as primeiras linhas; 30 sobra
LINHAS_ATE_O_CABECALHO = 30


def normalizar(texto: object) -> str:
    """Sem acento, minúsculo e sem espaço nas pontas: a planilha real tem
    `Estado ` e `PIS e COFINS Unit. ` com espaço no fim."""
    if texto is None:
        return ""
    sem_acento = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(sem_acento.lower().split())


@dataclass(frozen=True)
class Coluna:
    """Uma coluna esperada e os nomes que ela aceita, já normalizados. Nome
    terminado em `*` casa pelo começo — as colunas de custo têm a explicação
    entre parênteses, e ela muda de redação entre versões."""

    campo: str
    nomes: tuple[str, ...]
    obrigatoria: bool = True

    def casa(self, cabecalho: str) -> bool:
        return any(cabecalho.startswith(n[:-1]) if n.endswith("*") else cabecalho == n
                   for n in self.nomes)


COLUNAS_DA_BASE: tuple[Coluna, ...] = (
    Coluna("loja", ("nfcc_loj", "loja")),
    Coluna("cnpj", ("cnpj",), obrigatoria=False),
    Coluna("uf", ("estado", "uf")),
    Coluna("competencia", ("dt_ini", "competencia")),
    Coluna("item", ("cod_item",)),
    Coluna("descricao", ("descr_item",), obrigatoria=False),
    Coluna("quantidade", ("qtd_c",), obrigatoria=False),
    Coluna("valor_liquido", ("valor liquido",), obrigatoria=False),
    Coluna("status_custo", ("status custo",), obrigatoria=False),
    Coluna("base_da_analise", ("base da analise",)),
    Coluna("receita_bruta", ("receita bruta",)),
    Coluna("debito_icms", ("debito icms",)),
    Coluna("debito_ipi", ("debito ipi",), obrigatoria=False),
    Coluna("debito_pis_cofins", ("debito pis/cofins",)),
    Coluna("custo_liquido", ("custo liquido (sem*",)),
    Coluna("ipi_entrada", ("ipi na entrada*",)),
    Coluna("pis_cofins_sem_credito", ("pis/cofins sem credito*",)),
    Coluna("credito_icms", ("credito icms*",)),
    Coluna("credito_pis_cofins", ("credito pis/cofins*",)),
)

COLUNAS_DA_CLASSIFICACAO: tuple[Coluna, ...] = (
    Coluna("item", ("seqproduto", "cod_item", "codigo do item")),
    Coluna("descricao", ("desccompleta", "descricao", "descr_item"), obrigatoria=False),
    Coluna("ncm_original", ("ncm original",), obrigatoria=False),
    Coluna("ncm", ("ncm bms", "ncm revisada", "ncm")),
    # obrigatória para separar a revisão de SKU da aba "Base por SKU" da
    # planilha de resultado, que tem item, NCM e cClassTrib mas não o CST — e
    # que, lida como cadastro, entraria em dobro
    Coluna("cst", ("cst", "cst ibs/cbs")),
    Coluna("cclasstrib", ("cclasstrib",)),
    Coluna("categoria_seletivo", ("categoria is", "categoria do imposto seletivo",
                                  "imposto seletivo"), obrigatoria=False),
    Coluna("categoria_mercadologica", ("categoria mercadologica", "tipo mercadologico",
                                       "categoria"), obrigatoria=False),
)


class PlanilhaIlegivel(ValueError):
    """A planilha não tem o cabeçalho que se espera, ou não abre."""


@dataclass(frozen=True)
class Cabecalho:
    aba: str
    linha: int                      # 1-based, a do cabeçalho
    posicoes: dict[str, int]        # campo -> índice da coluna (0-based)
    faltam: tuple[str, ...]


def _casar(valores: tuple, colunas: tuple[Coluna, ...]) -> tuple[dict[str, int], tuple[str, ...]]:
    normalizados = [normalizar(v) for v in valores]
    posicoes: dict[str, int] = {}
    for coluna in colunas:
        for i, nome in enumerate(normalizados):
            if nome and coluna.casa(nome) and i not in posicoes.values():
                posicoes[coluna.campo] = i
                break
    faltam = tuple(c.campo for c in colunas if c.obrigatoria and c.campo not in posicoes)
    return posicoes, faltam


def achar_cabecalho(livro, colunas: tuple[Coluna, ...]) -> Cabecalho | None:
    """A primeira linha, em qualquer aba, onde todas as obrigatórias aparecem.
    Uma linha com mais da metade delas e alguma faltando é guardada como quase —
    é o que vira a mensagem de recusa, em vez de "não reconhecido"."""
    obrigatorias = sum(1 for c in colunas if c.obrigatoria)
    quase: Cabecalho | None = None
    for aba in livro.worksheets:
        for numero, valores in enumerate(
                aba.iter_rows(values_only=True, max_row=LINHAS_ATE_O_CABECALHO), start=1):
            posicoes, faltam = _casar(valores, colunas)
            if not faltam:
                return Cabecalho(aba.title, numero, posicoes, ())
            achadas = obrigatorias - len(faltam)
            if achadas * 2 > obrigatorias and (quase is None or len(faltam) < len(quase.faltam)):
                quase = Cabecalho(aba.title, numero, posicoes, faltam)
    return quase


def _abrir(caminho: str):
    from openpyxl import load_workbook  # noqa: PLC0415
    try:
        return load_workbook(caminho, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile) as erro:
        raise PlanilhaIlegivel(f"a planilha não abre: {erro}") from erro


@dataclass(frozen=True)
class Reconhecida:
    """O que a classificação do lote precisa saber de uma planilha da reforma."""

    especie: str                    # "base" ou "classificacao"
    cnpj: str | None
    competencia: date | None
    detalhe: str
    motivo: str = ""                # preenchido quando é quase: falta coluna


def reconhecer(caminho: str) -> Reconhecida | None:
    """Se a planilha é uma das duas da reforma. Lê só o cabeçalho e a primeira
    linha de dados. Nunca levanta: planilha que não abre não é da reforma."""
    try:
        livro = _abrir(caminho)
    except PlanilhaIlegivel:
        return None
    try:
        base = achar_cabecalho(livro, COLUNAS_DA_BASE)
        if base is not None and not base.faltam:
            primeira = _primeira_linha(livro, base)
            return Reconhecida("base", _cnpj(primeira, base), _competencia(primeira, base),
                               f"base custeada, aba {base.aba}")
        classificacao = achar_cabecalho(livro, COLUNAS_DA_CLASSIFICACAO)
        if classificacao is not None and not classificacao.faltam:
            return Reconhecida("classificacao", None, None,
                               f"classificação de itens, aba {classificacao.aba}")
        for quase, especie in ((base, "base custeada"), (classificacao, "classificação de itens")):
            if quase is not None:
                return Reconhecida("", None, None, "",
                                   motivo=f"parece {especie} da reforma, mas faltam as "
                                          f"colunas: {', '.join(quase.faltam)}")
        return None
    finally:
        livro.close()


def _primeira_linha(livro, cabecalho: Cabecalho) -> tuple:
    aba = livro[cabecalho.aba]
    for valores in aba.iter_rows(values_only=True, min_row=cabecalho.linha + 1,
                                 max_row=cabecalho.linha + 5):
        if any(v is not None for v in valores):
            return valores
    return ()


def _valor(linha: tuple, cabecalho: Cabecalho, campo: str):
    i = cabecalho.posicoes.get(campo)
    return linha[i] if i is not None and i < len(linha) else None


def _cnpj(linha: tuple, cabecalho: Cabecalho) -> str | None:
    digitos = "".join(c for c in str(_valor(linha, cabecalho, "cnpj") or "") if c.isdigit())
    return digitos if len(digitos) == 14 else None


def _competencia(linha: tuple, cabecalho: Cabecalho) -> date | None:
    valor = _valor(linha, cabecalho, "competencia")
    if isinstance(valor, datetime):
        return date(valor.year, valor.month, 1)
    if isinstance(valor, date):
        return date(valor.year, valor.month, 1)
    texto = str(valor or "").strip()
    for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d%m%Y", "%m/%Y"):
        try:
            d = datetime.strptime(texto[:10], formato)
            return date(d.year, d.month, 1)
        except ValueError:
            continue
    return None


# ---------------------------------------------------------------------------
# leitura completa
# ---------------------------------------------------------------------------

def chave_do_item(valor: object) -> str:
    """O código do item como texto: `1954343`, `1954343.0` e `"1954343"` são o
    mesmo SKU. O ERP grava número; a classificação às vezes grava texto."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    texto = str(valor).strip()
    if texto.endswith(".0") and texto[:-2].isdigit():
        return texto[:-2]
    return texto


def decimal(valor: object) -> Decimal:
    """Número da planilha em Decimal, pelo texto — `Decimal(0.1)` traria o erro
    binário junto. Vazio é zero: célula de tributo em branco é tributo nenhum."""
    if valor is None or valor == "":
        return Decimal(0)
    if isinstance(valor, (int, float)):
        return Decimal(repr(valor)) if isinstance(valor, float) else Decimal(valor)
    texto = str(valor).strip().replace(".", "").replace(",", ".") if "," in str(valor) else str(valor).strip()
    try:
        return Decimal(texto)
    except InvalidOperation:
        return Decimal(0)


@dataclass(frozen=True)
class LinhaDaBase:
    arquivo: str
    loja: str
    uf: str
    competencia: date | None
    item: str
    descricao: str
    quantidade: Decimal
    base_da_analise: str
    status_custo: str
    valor_liquido: Decimal
    receita_bruta: Decimal
    debito_icms: Decimal
    debito_ipi: Decimal
    debito_pis_cofins: Decimal
    custo_liquido: Decimal
    ipi_entrada: Decimal
    pis_cofins_sem_credito: Decimal
    credito_icms: Decimal
    credito_pis_cofins: Decimal
    tem_custo: bool                 # a célula do custo líquido veio preenchida


def ler_base(caminho: str, deve_parar: Callable[[], None] | None = None) -> Iterator[LinhaDaBase]:
    """As linhas de uma base custeada, em fluxo. Linha toda vazia é pulada."""
    livro = _abrir(caminho)
    try:
        cab = achar_cabecalho(livro, COLUNAS_DA_BASE)
        if cab is None or cab.faltam:
            raise PlanilhaIlegivel(
                f"{os.path.basename(caminho)}: não é a base custeada"
                + (f" — faltam {', '.join(cab.faltam)}" if cab else ""))
        nome = os.path.basename(caminho)
        for n, valores in enumerate(livro[cab.aba].iter_rows(values_only=True,
                                                             min_row=cab.linha + 1)):
            if deve_parar is not None and n % 5000 == 0:
                deve_parar()
            if not any(v is not None for v in valores):
                continue
            v = lambda campo: _valor(valores, cab, campo)  # noqa: E731
            yield LinhaDaBase(
                arquivo=nome,
                loja=chave_do_item(v("loja")),
                uf=str(v("uf") or "").strip().upper(),
                competencia=_competencia(valores, cab),
                item=chave_do_item(v("item")),
                descricao=str(v("descricao") or "").strip(),
                quantidade=decimal(v("quantidade")),
                base_da_analise=str(v("base_da_analise") or "").strip(),
                status_custo=str(v("status_custo") or "").strip(),
                valor_liquido=decimal(v("valor_liquido")),
                receita_bruta=decimal(v("receita_bruta")),
                debito_icms=decimal(v("debito_icms")),
                debito_ipi=decimal(v("debito_ipi")),
                debito_pis_cofins=decimal(v("debito_pis_cofins")),
                custo_liquido=decimal(v("custo_liquido")),
                ipi_entrada=decimal(v("ipi_entrada")),
                pis_cofins_sem_credito=decimal(v("pis_cofins_sem_credito")),
                credito_icms=decimal(v("credito_icms")),
                credito_pis_cofins=decimal(v("credito_pis_cofins")),
                tem_custo=v("custo_liquido") not in (None, ""),
            )
    finally:
        livro.close()


@dataclass(frozen=True)
class ItemClassificado:
    item: str
    descricao: str
    ncm_original: str
    ncm: str
    cst: str
    cclasstrib: str
    categoria_seletivo: str
    categoria_mercadologica: str


def ler_classificacao(caminho: str) -> Iterator[ItemClassificado]:
    livro = _abrir(caminho)
    try:
        cab = achar_cabecalho(livro, COLUNAS_DA_CLASSIFICACAO)
        if cab is None or cab.faltam:
            raise PlanilhaIlegivel(
                f"{os.path.basename(caminho)}: não é a classificação de itens"
                + (f" — faltam {', '.join(cab.faltam)}" if cab else ""))
        for valores in livro[cab.aba].iter_rows(values_only=True, min_row=cab.linha + 1):
            v = lambda campo: _valor(valores, cab, campo)  # noqa: E731
            item = chave_do_item(v("item"))
            if not item:
                continue
            yield ItemClassificado(
                item=item,
                descricao=str(v("descricao") or "").strip(),
                ncm_original=str(v("ncm_original") or "").strip(),
                ncm=str(v("ncm") or "").strip(),
                cst=str(v("cst") or "").strip(),
                cclasstrib=str(v("cclasstrib") or "").strip(),
                categoria_seletivo=str(v("categoria_seletivo") or "").strip(),
                categoria_mercadologica=str(v("categoria_mercadologica") or "").strip(),
            )
    finally:
        livro.close()
