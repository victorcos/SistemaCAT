"""As NF-e canceladas na SEFAZ que a EFD e o XML ainda trazem como válidas.

Na empresa 04, a RVZ achou notas canceladas na SEFAZ e ativas na escrituração
— o cliente não tinha a lista, e a ferramenta dela consultou a SEFAZ nota a
nota. Pedir ressarcimento sobre elas seria pedir sobre operação que não existe.

O sistema não consulta a SEFAZ: não guarda certificado de cliente. A cancelada
entra por duas vias (decisão do Victor, 16/09/2026):

* **evento de cancelamento** em XML (`procEventoNFe` com `tpEvento` 110111),
  que é o que se baixa junto com as notas do portal;
* **lista de chaves** em TXT, CSV ou planilha, com "cancel" no nome — a que a
  RVZ extraiu e a que o cliente mandou eram assim.

Sai daqui ``chaves_canceladas.parquet``; a consolidação da etapa 3 tira da
movimentação o que estiver nele, e conta.
"""

from __future__ import annotations

import os
import re
import zipfile
from dataclasses import dataclass, field

import pyarrow as pa
import pyarrow.parquet as pq

from cat.dominio.lote import abas_de_canceladas
from cat.log import obter_log

log = obter_log(__name__)

ARQUIVO_CHAVES_CANCELADAS = "chaves_canceladas.parquet"

ESQUEMA_CHAVES_CANCELADAS = pa.schema([
    ("chave", pa.string()),
    ("arquivo", pa.string()),
    ("origem", pa.string()),         # evento ou lista
])

_RE_CHAVE = re.compile(r"(?<!\d)(\d{44})(?!\d)")
_RE_CHAVE_DO_EVENTO = re.compile(rb"<chNFe>\s*(\d{44})\s*</chNFe>")
_RE_CANCELAMENTO = re.compile(rb"<tpEvento>\s*110111\s*</tpEvento>")
# a linha da lista que diz, ela mesma, que a nota não foi cancelada: o pedido de
# cancelamento rejeitado e o uso autorizado (na lista do cliente da empresa 04,
# 4 "Rejeição: Pedido de Cancelamento..." e 2 "Autorizado o uso da NF-e")
_RE_NAO_CANCELADA = re.compile(r"rejei|autorizado o uso|uso autorizado|denegad", re.IGNORECASE)
_RE_CANCELADA = re.compile(r"cancelamento\s+(autorizado|homologado)", re.IGNORECASE)


def linha_diz_que_nao_cancelou(textos: list[str]) -> bool:
    """A linha traz um retorno da SEFAZ que não é de cancelamento."""
    return (any(_RE_NAO_CANCELADA.search(t) for t in textos)
            and not any(_RE_CANCELADA.search(t) for t in textos))


@dataclass
class ProgressoDasCanceladas:
    arquivos: int = 0
    eventos: int = 0
    listas: int = 0
    chaves: int = 0
    ilegiveis: int = 0
    exemplos_ilegiveis: list[str] = field(default_factory=list)


def chaves_de_evento(conteudo: bytes) -> list[str]:
    """As chaves que um XML de evento cancela; nada se o evento não é cancelamento."""
    if not _RE_CANCELAMENTO.search(conteudo):
        return []
    return [m.decode("ascii") for m in _RE_CHAVE_DO_EVENTO.findall(conteudo)]


def chaves_de_texto(texto: str) -> list[str]:
    chaves = []
    for linha in texto.splitlines():
        if linha_diz_que_nao_cancelou([linha]):
            continue
        chaves += _RE_CHAVE.findall(linha)
    return chaves


def chaves_de_planilha(caminho: str) -> list[str]:
    """Toda chave de 44 dígitos escrita como texto nas abas de canceladas.

    Chave gravada como número o Excel já estragou (vira 3,52E+43) e não se
    recupera; fica de fora sem ser adivinhada. Aba de devoluções no mesmo
    arquivo não conta (`abas_de_canceladas`), nem a linha cujo retorno da SEFAZ
    diz que o cancelamento foi rejeitado ou que o uso está autorizado.
    """
    from openpyxl import load_workbook  # noqa: PLC0415

    chaves: list[str] = []
    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        validas = set(abas_de_canceladas(livro.sheetnames))
        ignoradas = [t for t in livro.sheetnames if t not in validas]
        if ignoradas:
            log.info("abas da planilha de canceladas que não listam canceladas",
                     extra={"arquivo": os.path.basename(caminho), "abas": ignoradas})
        fora = 0
        for aba in (a for a in livro.worksheets if a.title in validas):
            for linha in aba.iter_rows(values_only=True):
                textos = [v for v in linha if isinstance(v, str)]
                da_linha = [c for v in textos for c in _RE_CHAVE.findall(v)]
                if da_linha and linha_diz_que_nao_cancelou(textos):
                    fora += len(da_linha)
                    continue
                chaves += da_linha
        if fora:
            log.info("chaves da lista de canceladas cujo retorno da SEFAZ não é cancelamento",
                     extra={"arquivo": os.path.basename(caminho), "chaves": fora})
    finally:
        livro.close()
    return chaves


def ler_chaves_canceladas(caminhos: list[str], destino: str,
                          de_eventos: list[tuple[str, str]] = ()) -> ProgressoDasCanceladas:
    """Grava `chaves_canceladas.parquet` em `destino`, uma linha por chave.

    `de_eventos` são (chave, arquivo) de eventos que a leitura dos XML já achou
    — os que vêm dentro de zip, misturados com as notas.
    """
    os.makedirs(destino, exist_ok=True)
    progresso = ProgressoDasCanceladas(eventos=len(de_eventos))
    vistas: dict[str, tuple[str, str]] = {c: (a, "evento") for c, a in de_eventos}
    for caminho in caminhos:
        nome = os.path.basename(caminho)
        progresso.arquivos += 1
        try:
            if nome.lower().endswith(".xml"):
                with open(caminho, "rb") as f:
                    chaves, origem = chaves_de_evento(f.read()), "evento"
                progresso.eventos += 1
            elif nome.lower().endswith((".xlsx", ".xlsm")):
                chaves, origem = chaves_de_planilha(caminho), "lista"
                progresso.listas += 1
            else:
                with open(caminho, "rb") as f:
                    chaves, origem = chaves_de_texto(f.read().decode("latin-1")), "lista"
                progresso.listas += 1
        except (OSError, ValueError, KeyError, zipfile.BadZipFile) as erro:
            progresso.ilegiveis += 1
            if len(progresso.exemplos_ilegiveis) < 20:
                progresso.exemplos_ilegiveis.append(f"{nome}: {erro}")
            log.warning("arquivo de canceladas ilegível", extra={"arquivo": nome, "motivo": str(erro)})
            continue
        for chave in chaves:
            vistas.setdefault(chave, (nome, origem))
    pq.write_table(pa.Table.from_pylist(
        [{"chave": c, "arquivo": a, "origem": o} for c, (a, o) in sorted(vistas.items())],
        schema=ESQUEMA_CHAVES_CANCELADAS), os.path.join(destino, ARQUIVO_CHAVES_CANCELADAS))
    progresso.chaves = len(vistas)
    log.info("chaves canceladas lidas", extra={
        "arquivos": progresso.arquivos, "eventos": progresso.eventos, "listas": progresso.listas,
        "chaves": progresso.chaves, "ilegiveis": progresso.ilegiveis})
    return progresso
