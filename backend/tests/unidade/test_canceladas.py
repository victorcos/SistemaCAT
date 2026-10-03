"""As NF-e canceladas na SEFAZ: como entram no lote e como saem da movimentação.

A EFD é a de `test_movimentos`. A entrada com C170 está numa lista de chaves
(TXT) e a saída própria que o XML completa tem evento de cancelamento: as
duas saem da movimentação, ficam contadas, e o cupom continua.
"""

from __future__ import annotations

from datetime import date

import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from cat.dominio.lote import TipoDeArquivo
from cat.infraestrutura.analitico.canceladas import (
    ARQUIVO_CHAVES_CANCELADAS,
    chaves_de_evento,
    chaves_de_planilha,
    chaves_de_texto,
    ler_chaves_canceladas,
)
from cat.infraestrutura.analitico.itens_do_xml import extrair_itens_do_xml
from cat.infraestrutura.analitico.movimentacao import ARQUIVO_MOVIMENTOS, consolidar
from cat.infraestrutura.analitico.movimentos import extrair_movimentos
from cat.infraestrutura.arquivos.classificador import classificar
from tests.unidade.test_itens_do_xml import ICMS00, ICMS60, det, nfe
from tests.unidade.test_movimentos import (
    C100_ENTRADA, C100_SAIDA, C170_1, C170_2, C170_SEM_CADASTRO, C190_ENTRADA, C190_ENTRADA_60,
    C190_SAIDA_00, C190_SAIDA_60, C800, C850, CABECALHO, CHAVE_CUPOM, CHAVE_ENTRADA, CHAVE_SAIDA,
    CNPJ, CONV_A_FD, ITEM_A_V2, ITEM_B, escrever,
)

CHAVE_OUTRA = "35220744000001000454550010000123451000000018"


def evento(chave: str, tipo: str = "110111") -> bytes:
    return (f'<?xml version="1.0" encoding="UTF-8"?><procEventoNFe versao="1.00">'
            f"<evento><infEvento Id=\"ID{tipo}{chave}01\"><tpAmb>1</tpAmb><chNFe>{chave}</chNFe>"
            f"<tpEvento>{tipo}</tpEvento><detEvento><xJust>ERRO DE DIGITACAO</xJust></detEvento>"
            f"</infEvento></evento><retEvento><infEvento><cStat>135</cStat></infEvento></retEvento>"
            f"</procEventoNFe>").encode()


def planilha(caminho, linhas) -> str:
    livro = Workbook()
    aba = livro.active
    for linha in linhas:
        aba.append(linha)
    livro.save(caminho)
    return str(caminho)


class TestLeitura:
    def test_evento_de_cancelamento_da_a_chave(self):
        assert chaves_de_evento(evento(CHAVE_SAIDA)) == [CHAVE_SAIDA]

    def test_carta_de_correcao_nao_cancela(self):
        assert chaves_de_evento(evento(CHAVE_SAIDA, "110110")) == []

    def test_texto_pega_a_chave_e_ignora_numero_maior(self):
        texto = f"{CHAVE_ENTRADA}\r\nchave;{CHAVE_OUTRA};x\r\n{CHAVE_OUTRA}9\r\n"
        assert chaves_de_texto(texto) == [CHAVE_ENTRADA, CHAVE_OUTRA]

    def test_planilha_so_le_chave_escrita_como_texto(self, tmp_path):
        caminho = planilha(tmp_path / "canceladas.xlsx", [
            ["OBSERVACAO", "CHAVE NFE", "RETORNO SEFAZ"],
            ["", CHAVE_OUTRA, "Cancelamento autorizado"],
            ["inutilizada", None, "Inutilização de número homologado"],
            ["", 3.5220744000001e43, "Cancelamento autorizado"],
        ])
        assert chaves_de_planilha(caminho) == [CHAVE_OUTRA]

    def test_aba_de_devolucoes_no_mesmo_arquivo_nao_conta(self, tmp_path):
        """O relatório "NF-e Canceladas-Devoluções": uma aba de canceladas, outra de devoluções."""
        livro = Workbook()
        canceladas = livro.active
        canceladas.title = "NF-e Canceladas"
        canceladas.append(["Tipo de Operação", "Chave NFe"])
        canceladas.append(["Saída", CHAVE_OUTRA])
        devolucoes = livro.create_sheet("NF-e Devoluções")
        devolucoes.append(["ENTRADA", CHAVE_ENTRADA])
        caminho = str(tmp_path / "NF-e Canceladas-Devoluções.xlsx")
        livro.save(caminho)
        assert chaves_de_planilha(caminho) == [CHAVE_OUTRA]
        a = classificar(caminho)
        assert (a.tipo, a.detalhe) == (TipoDeArquivo.LISTA_DE_CANCELADAS, "1 chaves nas primeiras linhas")

    def test_linha_com_cancelamento_rejeitado_ou_uso_autorizado_nao_conta(self, tmp_path):
        caminho = planilha(tmp_path / "canceladas do cliente.xlsx", [
            ["OBSERVACAO", "CHAVE NFE", "RETORNO SEFAZ"],
            ["NF CANCELADA", CHAVE_OUTRA, "Cancelamento autorizado"],
            ["NF CANCELADA", CHAVE_ENTRADA, "Rejeição: Pedido de Cancelamento para NF-e com carta de correção"],
            ["NF CANCELADA", CHAVE_SAIDA, "Autorizado o uso da NF-e"],
        ])
        assert chaves_de_planilha(caminho) == [CHAVE_OUTRA]
        assert chaves_de_texto(f"{CHAVE_OUTRA};Cancelamento homologado fora de prazo\n{CHAVE_SAIDA};Autorizado o uso da NF-e\n") == [CHAVE_OUTRA]

    def test_abas_que_valem(self):
        from cat.dominio.lote import abas_de_canceladas
        assert abas_de_canceladas(["NF-e Canceladas", "NF-e Devoluções"]) == ["NF-e Canceladas"]
        assert abas_de_canceladas(["Planilha1"]) == ["Planilha1"]
        assert abas_de_canceladas(["Planilha1", "Devoluções"]) == ["Planilha1"]

    def test_uma_linha_por_chave_e_o_quebrado_contado(self, tmp_path):
        (tmp_path / "evento.xml").write_bytes(evento(CHAVE_SAIDA))
        (tmp_path / "canceladas.txt").write_text(f"{CHAVE_SAIDA}\n{CHAVE_ENTRADA}\n")
        (tmp_path / "canceladas.xlsx").write_bytes(b"PK\x03\x04 isto nao e planilha")
        destino = str(tmp_path / "saida")
        progresso = ler_chaves_canceladas(
            [str(tmp_path / n) for n in ("evento.xml", "canceladas.txt", "canceladas.xlsx")], destino)
        assert (progresso.arquivos, progresso.eventos, progresso.listas, progresso.chaves,
                progresso.ilegiveis) == (3, 1, 1, 2, 1)
        linhas = pq.read_table(f"{destino}/{ARQUIVO_CHAVES_CANCELADAS}").to_pylist()
        assert [(l["chave"], l["origem"]) for l in linhas] == [
            (CHAVE_ENTRADA, "lista"), (CHAVE_SAIDA, "evento")]


class TestClassificacao:
    def test_evento_de_cancelamento_alimenta(self, tmp_path):
        (tmp_path / "canc.xml").write_bytes(evento(CHAVE_OUTRA))
        a = classificar(str(tmp_path / "canc.xml"))
        assert a.tipo is TipoDeArquivo.XML_CANCELAMENTO
        assert (a.cnpj, a.competencia, a.alimenta_a_cat) == ("44000001000454", date(2022, 7, 1), True)

    def test_carta_de_correcao_continua_de_fora(self, tmp_path):
        (tmp_path / "cce.xml").write_bytes(evento(CHAVE_OUTRA, "110110"))
        assert classificar(str(tmp_path / "cce.xml")).tipo is TipoDeArquivo.XML_OUTRO

    def test_cancelamento_de_cte_nao_e_de_nota(self, tmp_path):
        cte = evento(CHAVE_OUTRA).replace(b"<chNFe>", b"<chCTe>").replace(b"</chNFe>", b"</chCTe>")
        (tmp_path / "cte.xml").write_bytes(cte.replace(b"procEventoNFe", b"procEventoCTe"))
        assert classificar(str(tmp_path / "cte.xml")).tipo is TipoDeArquivo.XML_OUTRO

    def test_lista_em_texto_precisa_do_nome(self, tmp_path):
        conteudo = f"{CHAVE_OUTRA}\r\n{CHAVE_ENTRADA}\r\n"
        lista = escrever(tmp_path, "NOTAS FISCAIS CANCELADAS - 072022 A 062024.txt", [conteudo])
        assert classificar(lista).tipo is TipoDeArquivo.LISTA_DE_CANCELADAS
        solta = escrever(tmp_path, "chaves.txt", [conteudo])
        assert classificar(solta).tipo is not TipoDeArquivo.LISTA_DE_CANCELADAS

    def test_planilha_com_chave_e_nome_entra(self, tmp_path):
        linhas = [["CHAVE NFE"], [CHAVE_OUTRA]]
        assert classificar(planilha(tmp_path / "NOTAS CANCELADAS ENVIADOS PELO CLIENTE.xlsx", linhas)
                           ).tipo is TipoDeArquivo.LISTA_DE_CANCELADAS
        assert classificar(planilha(tmp_path / "conferencia.xlsx", linhas)).tipo is TipoDeArquivo.DESCONHECIDO
        assert classificar(planilha(tmp_path / "canceladas_vazia.xlsx", [["CHAVE NFE"]])
                           ).tipo is TipoDeArquivo.DESCONHECIDO


@pytest.fixture
def pasta(tmp_path):
    efd = escrever(tmp_path, "efd_2021_05.txt", [
        CABECALHO, ITEM_A_V2, CONV_A_FD, ITEM_B,
        C100_ENTRADA, C170_1, C170_2, C170_SEM_CADASTRO, C190_ENTRADA, C190_ENTRADA_60,
        C100_SAIDA, C190_SAIDA_60, C190_SAIDA_00,
        C800, C850,
    ])
    xmls = tmp_path / "xml"
    xmls.mkdir()
    (xmls / "saida.xml").write_bytes(nfe(CHAVE_SAIDA, det(1, "1000144", "5405", "2", "250.00", ICMS60)
                                         + det(2, "X-MKT", "5102", "1", "50.00", ICMS00), emit=CNPJ))
    (xmls / "cupom.xml").write_bytes(
        (f'<?xml version="1.0"?><CFe><infCFe Id="CFe{CHAVE_CUPOM}"><ide><mod>59</mod>'
         f"<nserieSAT>353552</nserieSAT><nCFe>62537</nCFe><dEmi>20210503</dEmi></ide>"
         f"<emit><CNPJ>{CNPJ}</CNPJ></emit><dest/>"
         + det(1, "537861", "5405", "1.5", "36.20", ICMS60) + "</infCFe></CFe>").encode())
    (xmls / "cancelamento.xml").write_bytes(evento(CHAVE_SAIDA))
    (tmp_path / "canceladas.txt").write_text(f"{CHAVE_ENTRADA}\n{CHAVE_OUTRA}\n")
    destino = str(tmp_path / "saida")
    extrair_movimentos([efd], destino)
    extrair_itens_do_xml([str(xmls / "saida.xml"), str(xmls / "cupom.xml")], destino)
    ler_chaves_canceladas([str(xmls / "cancelamento.xml"), str(tmp_path / "canceladas.txt")], destino)
    resumo = consolidar(destino, None)
    linhas = pq.read_table(f"{destino}/{ARQUIVO_MOVIMENTOS}").to_pylist()
    return resumo, linhas


class TestConsolidacao:
    def test_cancelada_sai_da_movimentacao(self, pasta):
        _, linhas = pasta
        assert {l["chave"] for l in linhas} == {CHAVE_CUPOM}

    def test_e_fica_contada(self, pasta):
        resumo, _ = pasta
        # a entrada com 3 C170 e a saída com os 2 itens do XML; a chave que a
        # EFD não tem não conta como documento
        assert (resumo.chaves_canceladas, resumo.documentos_cancelados_na_sefaz,
                resumo.movimentos_cancelados) == (3, 2, 5)
        assert (resumo.saidas_completadas_pelo_xml, resumo.movimentos_do_xml) == (1, 1)
        assert any("cancelados na SEFAZ" in a for a in resumo.avisos)

    def test_sem_lista_nada_muda(self, tmp_path):
        efd = escrever(tmp_path, "efd.txt", [CABECALHO, ITEM_A_V2, CONV_A_FD, ITEM_B,
                                              C100_ENTRADA, C170_1, C190_ENTRADA])
        destino = str(tmp_path / "saida")
        extrair_movimentos([efd], destino)
        resumo = consolidar(destino, None)
        assert (resumo.chaves_canceladas, resumo.documentos_cancelados_na_sefaz) == (0, 0)
        assert resumo.movimentos == 1
