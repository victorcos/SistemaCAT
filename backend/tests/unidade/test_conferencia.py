"""Conferência de documentos: leitura do C100/C800, confronto e planilhas.

As duas linhas de registro abaixo foram **copiadas de arquivos reais** da
empresa 12. Foi assim que os leiautes deste projeto se confirmaram,
e não pelo manual: já aconteceu de arquivo real não bater com o leiaute
publicado, e um campo deslocado aqui trocaria chave de nota por número de nota
em milhões de linhas.
"""

import csv
import re
from datetime import date
from decimal import Decimal

import duckdb
import pytest

from cat.dominio.icms.cat42.conferencia import ResumoDaConferencia
from cat.dominio.sped.fiscais import Emitente, Operacao, ler_documento
from cat.infraestrutura.analitico.confronto import confrontar
from cat.infraestrutura.analitico.extracao import extrair_efd, extrair_pasta
from cat.infraestrutura.planilhas.conferencia import (
    gerar_conferidas,
    gerar_nao_escrituradas,
    gerar_sem_documento,
)

# real: estabelecimento do Paraná, competência 05/2021
C100 = (
    "|C100|0|0|C10252525|55|00|001|44623"
    "|41210544000002000237550010000446231411953286"
    "|01052021|01052021|8,49|2|0,49|0|8,98|9|0|0|0|0|0|0|0|0|0|0|0|0|"
)
# real: CF-e-SAT de um estabelecimento de São Paulo
C800 = (
    "|C800|59|00|62537|01072021|2,59|0|0||353552"
    "|35210744000002005468590003535520625376456957|0|2,59|0|0|0|0|"
)
CABECALHO = (
    "|0000|015|0|01052021|31052021|Empresa 12 Distribuicao Ltda"
    "|44000002000237||PR|9000000000|4106407||||"
)

CHAVE_C100 = "41210544000002000237550010000446231411953286"
CHAVE_C800 = "35210744000002005468590003535520625376456957"

NFE = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<nfeProc versao="4.00"><NFe><infNFe Id="NFe'
    + CHAVE_C100
    + '"><ide><dhEmi>2021-05-01T10:00:00-03:00</dhEmi></ide>'
    "<emit><CNPJ>44000002000237</CNPJ></emit></infNFe></NFe></nfeProc>"
)


class TestLerC100:
    def test_campos_na_posicao_certa(self):
        d = ler_documento(C100)
        assert d is not None
        assert d.chave == CHAVE_C100
        assert d.modelo == "55"
        assert d.situacao == "00"
        assert d.serie == "001"
        assert d.numero == "44623"
        assert d.data == date(2021, 5, 1)
        assert d.valor == Decimal("8.49")
        assert d.participante == "C10252525"

    def test_indicadores_de_operacao_e_emitente(self):
        d = ler_documento(C100)
        assert d.operacao is Operacao.ENTRADA      # IND_OPER = 0
        assert d.emitente is Emitente.PROPRIA      # IND_EMIT = 0

    def test_saida_de_terceiros(self):
        d = ler_documento(C100.replace("|C100|0|0|", "|C100|1|1|", 1))
        assert d.operacao is Operacao.SAIDA
        assert d.emitente is Emitente.TERCEIROS


class TestLerC800:
    def test_campos_na_posicao_certa(self):
        d = ler_documento(C800)
        assert d is not None
        assert d.chave == CHAVE_C800
        assert d.modelo == "59"
        assert d.numero == "62537"
        assert d.data == date(2021, 7, 1)
        assert d.valor == Decimal("2.59")

    def test_cupom_e_sempre_saida_propria(self):
        d = ler_documento(C800)
        assert d.operacao is Operacao.SAIDA
        assert d.emitente is Emitente.PROPRIA


class TestSituacao:
    @pytest.mark.parametrize(("codigo", "vale"), [
        ("00", True), ("01", True), ("06", True), ("07", True), ("08", True),
        ("02", False), ("03", False), ("04", False), ("05", False),
    ])
    def test_o_que_entra_na_apuracao(self, codigo, vale):
        d = ler_documento(C100.replace("|55|00|", f"|55|{codigo}|", 1))
        assert d.vale_na_apuracao is vale

    def test_cancelada_nao_se_cobra(self):
        # não existe documento a pedir para nota cancelada; mandar isso numa
        # planilha de cobrança queima a conversa com o cliente
        d = ler_documento(C100.replace("|55|00|", "|55|02|", 1))
        assert not d.pode_ser_cobrado

    def test_sem_chave_e_reconhecido(self):
        d = ler_documento(C100.replace(CHAVE_C100, "", 1))
        assert d.sem_chave


class TestOutrasLinhas:
    @pytest.mark.parametrize("linha", [
        "|C190|000|1102|18,00|100,00|18,00|0|0|0|0|0||",
        "|C170|1|116785|Laranja|200|KG|417,58|",
        "|0000|015|0|01052021|31052021|Empresa|44000002000237||PR|1|1||||",
    ])
    def test_nao_sao_documento(self, linha):
        assert ler_documento(linha) is None


# ---------------------------------------------------------------------------
# extração e confronto
# ---------------------------------------------------------------------------
def escrever(pasta, nome, linhas, codificacao="latin-1"):
    caminho = pasta / nome
    caminho.write_bytes(("\r\n".join(linhas) + "\r\n").encode(codificacao))
    return str(caminho)


@pytest.fixture
def efd(tmp_path):
    return escrever(tmp_path, "efd.txt", [
        CABECALHO,
        C100,
        C800,
        # cancelada: fica na lista, marcada — não some do controle
        C100.replace("|55|00|001|44623|" + CHAVE_C100,
                     "|55|02|001|44624|" + CHAVE_C100[:-1] + "0"),
        # sem chave: não dá para confrontar, mas entra marcada
        C100.replace(CHAVE_C100, "").replace("|44623|", "|44625|"),
        "|C190|000|1102|18,00|100,00|18,00|0|0|0|0|0||",
    ])


class TestExtrair:
    def test_le_so_c100_e_c800(self, efd, tmp_path):
        destino = str(tmp_path / "efd.parquet")
        progresso = extrair_efd([efd], destino)
        assert progresso.documentos == 4          # o C190 e o 0000 não contam

        con = duckdb.connect()
        linhas = con.execute(
            "SELECT chave, modelo, cnpj, competencia FROM read_parquet(?) "
            "ORDER BY modelo", [destino]).fetchall()
        con.close()
        # o CNPJ e a competência vêm do registro 0000 do próprio arquivo
        assert all(l[2] == "44000002000237" for l in linhas)
        assert all(l[3] == date(2021, 5, 1) for l in linhas)
        assert {l[1] for l in linhas} == {"55", "59"}

    def test_arquivo_ilegivel_nao_derruba_os_outros(self, efd, tmp_path):
        progresso = extrair_efd([efd, str(tmp_path / "nao_existe.txt")],
                                str(tmp_path / "efd.parquet"))
        assert progresso.documentos == 4
        assert len(progresso.recusados) == 1

    def test_chave_do_xml_pelo_conteudo(self, tmp_path):
        caminho = tmp_path / "nota.xml"
        caminho.write_text(NFE, encoding="utf-8")
        destino = str(tmp_path / "pasta.parquet")
        progresso = extrair_pasta([str(caminho)], [], destino)
        assert progresso.documentos == 1

        con = duckdb.connect()
        assert con.execute("SELECT chave, origem FROM read_parquet(?)",
                           [destino]).fetchone() == (CHAVE_C100, "xml")
        con.close()

    def test_chave_pelo_nome_quando_o_conteudo_nao_tem(self, tmp_path):
        # muitos exportadores nomeiam o arquivo pela própria chave
        caminho = tmp_path / f"{CHAVE_C800}-procNFe.xml"
        caminho.write_text('<?xml version="1.0"?><outro/>', encoding="utf-8")
        progresso = extrair_pasta([str(caminho)], [],
                                  str(tmp_path / "pasta.parquet"))
        assert progresso.documentos == 1

    def test_xml_sem_chave_nenhuma_e_recusado(self, tmp_path):
        caminho = tmp_path / "evento.xml"
        caminho.write_text('<?xml version="1.0"?><evento/>', encoding="utf-8")
        progresso = extrair_pasta([str(caminho)], [],
                                  str(tmp_path / "pasta.parquet"))
        assert progresso.documentos == 0
        assert progresso.recusados


@pytest.fixture
def confronto(efd, tmp_path):
    """A EFD tem 4 documentos; a pasta traz um deles e um que não existe lá."""
    caminho_efd = str(tmp_path / "efd.parquet")
    extrair_efd([efd], caminho_efd)

    xml = tmp_path / "nota.xml"
    xml.write_text(NFE, encoding="utf-8")
    intruso = tmp_path / f"{'7' * 44}-nfe.xml"
    intruso.write_text('<?xml version="1.0"?><nfeProc/>', encoding="utf-8")

    caminho_pasta = str(tmp_path / "pasta.parquet")
    extrair_pasta([str(xml), str(intruso)], [], caminho_pasta)

    destino = str(tmp_path / "saida")
    return confrontar(caminho_efd, caminho_pasta, destino), destino


class TestConfrontar:
    def test_a_conta_fecha(self, confronto):
        resumo, _ = confronto
        assert resumo.conferidos + resumo.sem_documento == resumo.escriturados

    def test_o_que_tem_documento(self, confronto):
        resumo, _ = confronto
        assert resumo.conferidos == 1

    def test_a_lista_positiva_traz_a_nota_e_de_onde_veio_o_documento(
        self, confronto
    ):
        # "possui nota e está na EFD — onde vejo?" Aqui: uma linha por
        # chave, com o arquivo do documento que casou
        resumo, destino = confronto
        con = duckdb.connect()
        linhas = con.execute(
            "SELECT chave, origem, arquivo_do_documento, ocorrencias "
            f"FROM read_parquet('{destino}/conferidos.parquet')").fetchall()
        con.close()
        assert len(linhas) == resumo.conferidos == 1
        chave, origem, arquivo, ocorrencias = linhas[0]
        assert chave == CHAVE_C100
        assert origem == "xml"
        assert arquivo.endswith("nota.xml")
        assert ocorrencias == 1

    def test_as_duas_listas_particionam_a_efd(self, confronto):
        # cada documento escriturado está em exatamente uma das listas
        resumo, destino = confronto
        con = duckdb.connect()
        positivas = con.execute(
            f"SELECT count(*) FROM read_parquet('{destino}/conferidos.parquet')"
        ).fetchone()[0]
        pendentes = con.execute(
            f"SELECT count(*) FROM read_parquet('{destino}/sem_documento.parquet')"
        ).fetchone()[0]
        con.close()
        assert positivas + pendentes == resumo.escriturados == 4

    def test_o_que_esta_na_pasta_e_nao_na_efd(self, confronto):
        resumo, _ = confronto
        assert resumo.nao_escrituradas == 1

    def test_documento_sem_chave_entra_na_lista_marcado(self, confronto):
        # não dá para confrontar sem chave, mas excluir o tiraria do controle
        # para sempre: na rodada seguinte ninguém voltaria a olhá-lo
        resumo, destino = confronto
        assert resumo.sem_chave_na_efd == 1
        assert resumo.escriturados == 4
        assert _classificacoes(destino)["sem_chave"] == 1

    def test_cancelada_fica_na_lista_marcada(self, confronto):
        resumo, destino = confronto
        assert resumo.sem_documento == 3
        # só uma espera documento; as outras duas ficam marcadas do que são
        assert resumo.sem_documento_cobravel == 1
        classes = _classificacoes(destino)
        assert classes["a_cobrar"] == 1
        assert classes["sem_documento_a_pedir"] == 1

    def test_cobertura(self, confronto):
        resumo, _ = confronto
        assert resumo.cobertura == pytest.approx(1 / 4)

    def test_recorte_por_modelo_traz_o_codigo(self, confronto):
        resumo, _ = confronto
        codigos = {f.codigo for f in resumo.por_modelo}
        assert codigos == {"55", "59"}

    def test_a_lista_tem_uma_linha_por_chave(self, confronto):
        # a pergunta que veio da operação: "as notas para cobrar não estão
        # duplicadas?" Não estão, por construção — e a guarda mede isso
        resumo, destino = confronto
        con = duckdb.connect()
        com_chave, distintas = con.execute(
            f"SELECT count(*), count(DISTINCT chave) "
            f"FROM read_parquet('{destino}/sem_documento.parquet') WHERE tem_chave"
        ).fetchone()
        assert com_chave == distintas
        assert resumo.chaves_repetidas_na_lista == 0
        assert not any("chave repetida" in a for a in resumo.avisos)


class TestNotaEmDuasFiliaisComDocumento:
    """A mesma chave em dois C100 (emitida por uma filial, recebida por outra)
    e o XML na pasta: é UMA nota conferida, com duas ocorrências — o mesmo
    tratamento que a lista de pendências dá."""

    def test_conta_uma_vez_e_mostra_as_ocorrencias(self, efd, tmp_path):
        import shutil  # noqa: PLC0415

        outra_filial = str(tmp_path / "efd_filial.txt")
        shutil.copy(efd, outra_filial)
        caminho_efd = str(tmp_path / "efd.parquet")
        extrair_efd([efd, outra_filial], caminho_efd)

        xml = tmp_path / "nota.xml"
        xml.write_text(NFE, encoding="utf-8")
        caminho_pasta = str(tmp_path / "pasta.parquet")
        extrair_pasta([str(xml)], [], caminho_pasta)

        destino = str(tmp_path / "saida")
        resumo = confrontar(caminho_efd, caminho_pasta, destino)
        assert resumo.conferidos == 1
        assert resumo.escriturados == 4          # 8 linhas, 4 documentos

        con = duckdb.connect()
        assert con.execute(
            f"SELECT ocorrencias FROM read_parquet('{destino}/conferidos.parquet')"
        ).fetchall() == [(2,)]
        con.close()


class TestOndeFicaORascunho:
    """O XML que o xlsxwriter despeja enquanto monta a planilha.

    Com `constant_memory` cada linha vai para disco na hora. Para 8,7 milhões
    de linhas isso deu 7 a 9 GB de rascunho, e na pasta temporária do sistema
    — que nesta casa fica no C:, o menor disco da máquina — o download enchia
    o disco antes de terminar. O rascunho passa a cair ao lado do destino.
    """

    def test_nao_escreve_na_pasta_temporaria_do_sistema(
        self, confronto, tmp_path, monkeypatch
    ):
        """Se voltar a usar o TEMP do sistema, este teste denuncia."""
        proibida = tmp_path / "temp_do_sistema"
        proibida.mkdir()
        for variavel in ("TMPDIR", "TEMP", "TMP"):
            monkeypatch.setenv(variavel, str(proibida))

        _, destino = confronto
        saida = tmp_path / "planilha" / "p.xlsx"
        saida.parent.mkdir()
        gerar_sem_documento(f"{destino}/sem_documento.parquet", str(saida))

        assert saida.is_file()
        assert list(proibida.iterdir()) == [], "escreveu rascunho no TEMP do sistema"

    def test_a_pasta_do_destino_e_criada_se_faltar(self, confronto, tmp_path):
        """O destino pode apontar para pasta que ainda não existe."""
        _, destino = confronto
        saida = tmp_path / "que" / "nao" / "existe" / "p.xlsx"
        linhas = gerar_sem_documento(f"{destino}/sem_documento.parquet", str(saida))
        assert linhas == 3
        assert saida.is_file()

    def test_o_csv_nao_precisa_de_rascunho(self, confronto, tmp_path, monkeypatch):
        """CSV escreve direto no destino: é por isso que ele não custa disco
        de rascunho, e é a saída certa para lista grande."""
        proibida = tmp_path / "temp_do_sistema"
        proibida.mkdir()
        for variavel in ("TMPDIR", "TEMP", "TMP"):
            monkeypatch.setenv(variavel, str(proibida))
        _, destino = confronto
        saida = tmp_path / "p.csv"
        gerar_sem_documento(f"{destino}/sem_documento.parquet", str(saida),
                            formato="csv")
        assert saida.is_file()
        assert list(proibida.iterdir()) == []


class TestCsv:
    """A outra saída da mesma lista.

    O que precisa valer: mesmo conteúdo do xlsx, e um arquivo que o Excel em
    português e o DuckDB leiam sem tratamento.
    """

    def _ler(self, caminho: str) -> list[list[str]]:
        # utf-8-sig porque o arquivo tem BOM; sem ele o Excel lê como ANSI e
        # todo acento vira lixo
        with open(caminho, encoding="utf-8-sig", newline="") as f:
            return list(csv.reader(f, delimiter=";"))

    def test_tem_as_mesmas_linhas_que_o_xlsx(self, confronto, tmp_path):
        _, destino = confronto
        origem = f"{destino}/sem_documento.parquet"
        no_xlsx = gerar_sem_documento(origem, str(tmp_path / "p.xlsx"))
        no_csv = gerar_sem_documento(origem, str(tmp_path / "p.csv"),
                                     formato="csv")
        assert no_csv == no_xlsx == 3
        # cabeçalho + as três linhas
        assert len(self._ler(str(tmp_path / "p.csv"))) == 4

    def test_o_filtro_vale_igual_nos_dois_formatos(self, confronto, tmp_path):
        """Filtro divergente daria dois totais para a mesma cobrança."""
        _, destino = confronto
        origem = f"{destino}/sem_documento.parquet"
        so = frozenset({"a_cobrar"})
        assert gerar_sem_documento(origem, str(tmp_path / "f.csv"),
                                   classificacoes=so, formato="csv") == 1
        assert gerar_sem_documento(origem, str(tmp_path / "f.xlsx"),
                                   classificacoes=so) == 1

    def test_tem_bom_e_ponto_e_virgula(self, confronto, tmp_path):
        """Sem BOM o Excel come o acento; com vírgula de separador, todo
        valor com centavo quebraria a coluna."""
        _, destino = confronto
        alvo = str(tmp_path / "p.csv")
        gerar_sem_documento(f"{destino}/sem_documento.parquet", alvo,
                            formato="csv")
        with open(alvo, "rb") as f:
            comeco = f.read(200)
        assert comeco[:3] == b"\xef\xbb\xbf"
        assert b";" in comeco

    def test_a_chave_sai_com_os_44_digitos(self, confronto, tmp_path):
        """O CSV não finge tipo para agradar o Excel.

        É o defeito que este projeto diagnosticou num relatório de cliente:
        44 dígitos abertos no Excel viram "4,12105E+43" e não voltam. Quem
        precisa do Excel baixa o xlsx, que é imune.
        """
        _, destino = confronto
        alvo = str(tmp_path / "p.csv")
        gerar_sem_documento(f"{destino}/sem_documento.parquet", alvo,
                            formato="csv")
        cabecalho, *linhas = self._ler(alvo)
        onde = cabecalho.index("Chave de acesso")
        chaves = [l[onde] for l in linhas if l[onde]]
        assert chaves
        for chave in chaves:
            assert len(chave) == 44 and chave.isdigit(), chave

    def test_numero_com_virgula_decimal_e_data_no_formato_daqui(
        self, confronto, tmp_path
    ):
        _, destino = confronto
        alvo = str(tmp_path / "p.csv")
        gerar_sem_documento(f"{destino}/sem_documento.parquet", alvo,
                            formato="csv")
        cabecalho, *linhas = self._ler(alvo)
        valor = linhas[0][cabecalho.index("Valor do documento")]
        assert "," in valor and "." not in valor
        emissao = linhas[0][cabecalho.index("Emissão")]
        assert re.fullmatch(r"\d{2}/\d{2}/\d{4}", emissao), emissao

    def test_rotulo_traduzido_igual_ao_xlsx(self, confronto, tmp_path):
        """A tradução mora no gerador, não no formato."""
        _, destino = confronto
        alvo = str(tmp_path / "p.csv")
        gerar_sem_documento(f"{destino}/sem_documento.parquet", alvo,
                            formato="csv")
        cabecalho, *linhas = self._ler(alvo)
        operacoes = {l[cabecalho.index("Operação")] for l in linhas}
        assert operacoes <= {"Entrada", "Saída"}, operacoes

    def test_formato_desconhecido_e_recusado(self, confronto, tmp_path):
        _, destino = confronto
        with pytest.raises(ValueError, match="formato desconhecido"):
            gerar_sem_documento(f"{destino}/sem_documento.parquet",
                                str(tmp_path / "p.ods"), formato="ods")


class TestPlanilhas:
    def test_sai_inteira_sem_filtro(self, confronto, tmp_path):
        # nada é excluído: cancelada e sem chave vão junto, marcadas
        _, destino = confronto
        linhas = gerar_sem_documento(f"{destino}/sem_documento.parquet",
                                     str(tmp_path / "pendencias.xlsx"))
        assert linhas == 3

    def test_filtro_por_classificacao(self, confronto, tmp_path):
        _, destino = confronto
        so_cobrar = gerar_sem_documento(
            f"{destino}/sem_documento.parquet", str(tmp_path / "cobrar.xlsx"),
            classificacoes=frozenset({"a_cobrar"}))
        assert so_cobrar == 1

    def test_filtro_por_modelo(self, confronto, tmp_path):
        # numa base real, 307.319 de 321.337 documentos eram NFC-e; sem filtro
        # a planilha de cobrança fica inútil pelo volume
        _, destino = confronto
        so_nfe = gerar_sem_documento(f"{destino}/sem_documento.parquet",
                                     str(tmp_path / "so55.xlsx"),
                                     modelos=frozenset({"55"}))
        so_cupom = gerar_sem_documento(f"{destino}/sem_documento.parquet",
                                       str(tmp_path / "so59.xlsx"),
                                       modelos=frozenset({"59"}))
        assert so_nfe + so_cupom == 3

    def test_nao_escrituradas(self, confronto, tmp_path):
        _, destino = confronto
        linhas = gerar_nao_escrituradas(f"{destino}/nao_escrituradas.parquet",
                                        str(tmp_path / "fora.xlsx"))
        assert linhas == 1

    def test_conferidas_diz_de_onde_veio_o_documento(self, confronto, tmp_path):
        import openpyxl  # noqa: PLC0415

        _, destino = confronto
        caminho = str(tmp_path / "conferidas.xlsx")
        assert gerar_conferidas(f"{destino}/conferidos.parquet", caminho) == 1
        aba = openpyxl.load_workbook(caminho).active
        colunas = [c.value for c in aba[1]]
        origem = aba.cell(row=2, column=colunas.index("Origem do documento") + 1)
        assert origem.value == "XML"
        assert "Classificação" not in colunas       # quem casou, casou

    def test_chave_vai_como_texto(self, confronto, tmp_path):
        # chave de 44 dígitos como número vira notação científica: parece certa
        # e não é, que é pior do que dar erro
        import openpyxl  # noqa: PLC0415

        _, destino = confronto
        caminho = str(tmp_path / "cobrar.xlsx")
        gerar_sem_documento(f"{destino}/sem_documento.parquet", caminho)
        aba = openpyxl.load_workbook(caminho).active
        colunas = [c.value for c in aba[1]]
        celula = aba.cell(row=2, column=colunas.index("Chave de acesso") + 1)
        assert isinstance(celula.value, str)
        assert len(celula.value) == 44


def _classificacoes(destino: str) -> dict[str, int]:
    con = duckdb.connect()
    try:
        return dict(con.execute(
            "SELECT classificacao, count(*) FROM read_parquet(?) GROUP BY 1",
            [f"{destino}/sem_documento.parquet"]).fetchall())
    finally:
        con.close()


class TestSegundaRodada:
    """O trabalho não termina na primeira conferência.

    O cliente manda o que faltava e a conferência roda de novo. Sem comparar
    com a rodada anterior, a segunda só diz "ainda faltam N" e ninguém sabe se
    andou — que foi o que motivou não excluir nada da lista.
    """

    def test_diz_quantas_o_cliente_resolveu(self, efd, tmp_path):
        caminho_efd = str(tmp_path / "efd.parquet")
        extrair_efd([efd], caminho_efd)

        primeira = str(tmp_path / "r1")
        vazia = str(tmp_path / "vazia.parquet")
        extrair_pasta([], [], vazia)
        antes = confrontar(caminho_efd, vazia, primeira)

        # agora o cliente mandou o XML de uma das notas
        xml = tmp_path / "nota.xml"
        xml.write_text(NFE, encoding="utf-8")
        com_xml = str(tmp_path / "com_xml.parquet")
        extrair_pasta([str(xml)], [], com_xml)

        depois = confrontar(caminho_efd, com_xml, str(tmp_path / "r2"),
                            anterior=f"{primeira}/sem_documento.parquet")

        assert depois.comparou
        assert depois.pendencias_resolvidas == 1
        assert depois.pendencias_novas == 0
        assert depois.sem_documento == antes.sem_documento - 1
        assert "1 resolvida(s)" in depois.andou

    def test_sem_rodada_anterior_nao_compara(self, efd, tmp_path):
        caminho_efd = str(tmp_path / "efd.parquet")
        extrair_efd([efd], caminho_efd)
        vazia = str(tmp_path / "vazia.parquet")
        extrair_pasta([], [], vazia)
        r = confrontar(caminho_efd, vazia, str(tmp_path / "r1"),
                       anterior=str(tmp_path / "nao_existe.parquet"))
        assert not r.comparou
        assert r.andou == ""


class TestResumo:
    def test_sem_origem_avisa_que_nao_diz_nada(self):
        r = ResumoDaConferencia(escriturados=10)
        assert any("não diz nada" in a for a in r.avisos)

    def test_aviso_nao_come_a_pontuacao_do_texto(self):
        # o separador de milhar é aplicado ao número, não à frase
        r = ResumoDaConferencia(escriturados=1, sem_chave_na_efd=1234)
        aviso = next(a for a in r.avisos if "1.234" in a)
        assert "modelo 1 ou cupom antigo." in aviso

    def test_cobertura_sem_documento_nenhum(self):
        assert ResumoDaConferencia().cobertura == 0.0


class TestFiliaisTrocadas:
    """O erro de uso mais provável, e o mais confuso quando não é explicado.

    Importar a EFD de uma filial e os XML de outra dá cobertura 0% e uma lista
    de milhares de pendências que parecem falta de documento — quando na
    verdade nada podia casar. Aconteceu num uso real: EFD do 0001-89 contra XML
    do 0004-21, 10.605 pendências e nenhum documento conferido.
    """

    def test_avisa_quando_os_estabelecimentos_nao_se_cruzam(self):
        r = ResumoDaConferencia(
            escriturados=10605, conferidos=0, documentos_na_pasta=882,
            estabelecimentos_da_efd=["44000001000101"],
            emitentes_na_pasta=["44000001000454"],
        )
        aviso = next(a for a in r.avisos if "filiais diferentes" in a)
        assert "44000001000101" in aviso
        assert "44000001000454" in aviso

    def test_nao_avisa_quando_ha_intersecao(self):
        r = ResumoDaConferencia(
            escriturados=10, conferidos=0, documentos_na_pasta=5,
            estabelecimentos_da_efd=["44000001000101"],
            emitentes_na_pasta=["44000001000101", "99999999000191"],
        )
        assert not any("filiais diferentes" in a for a in r.avisos)

    def test_nao_avisa_quando_algo_casou(self):
        # com pelo menos um casamento, a causa não é filial trocada
        r = ResumoDaConferencia(
            escriturados=10, conferidos=1, documentos_na_pasta=5,
            estabelecimentos_da_efd=["44000001000101"],
            emitentes_na_pasta=["44000001000454"],
        )
        assert not any("filiais diferentes" in a for a in r.avisos)


class TestDuplicidadeDeArquivo:
    """`ocorrencias` na planilha, e aviso quando a repetição denuncia arquivo em dobro.

    O corte é 5%. Numa base real de 37,9 milhões, 2,2% das chaves se repetiam
    legitimamente (a mesma nota escriturada em duas filiais); o mesmo lote
    importado duas vezes daria perto de 100%.
    """

    def test_acima_do_corte_avisa(self):
        r = ResumoDaConferencia(escriturados=100, sem_documento=100,
                                pendencias_repetidas=40)
        aviso = next(a for a in r.avisos if "arquivo importado em dobro" in a)
        assert "40%" in aviso

    def test_repeticao_legitima_nao_avisa(self):
        # 2% é filial repetida, não arquivo em dobro
        r = ResumoDaConferencia(escriturados=100, sem_documento=100,
                                pendencias_repetidas=2)
        assert not any("em dobro" in a for a in r.avisos)

    def test_chave_repetida_na_lista_e_defeito_e_avisa(self):
        # nunca deve acontecer; se acontecer, a tela diz para não cobrar
        r = ResumoDaConferencia(escriturados=10, sem_documento=10,
                                chaves_repetidas_na_lista=3)
        aviso = next(a for a in r.avisos if "chave repetida" in a)
        assert "3 linha(s) a mais" in aviso
        assert "Não cobre" in aviso

    def test_sem_pendencia_nao_divide_por_zero(self):
        assert ResumoDaConferencia(escriturados=5, sem_documento=0).avisos == [
            a for a in ResumoDaConferencia(escriturados=5, sem_documento=0).avisos
        ]

    def test_a_planilha_mostra_as_ocorrencias(self, confronto, tmp_path):
        import openpyxl  # noqa: PLC0415

        _, destino = confronto
        caminho = str(tmp_path / "pendencias.xlsx")
        gerar_sem_documento(f"{destino}/sem_documento.parquet", caminho)
        aba = openpyxl.load_workbook(caminho).active
        cabecalho = [c.value for c in aba[1]]
        assert "Ocorrências na EFD" in cabecalho
        coluna = cabecalho.index("Ocorrências na EFD") + 1
        # todas as pendências da fixture são chaves únicas: 1 ocorrência
        assert aba.cell(row=2, column=coluna).value == 1
