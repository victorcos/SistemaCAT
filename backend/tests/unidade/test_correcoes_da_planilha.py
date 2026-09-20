"""A Ficha 3 sobe editada e o sistema diz o que mudou.

O caminho inteiro num teste só: gera a planilha do razão, mexe nas células como
uma pessoa mexeria no Excel, sobe de volta e confere o diff. É o que separa
"subir planilha" de "gravar o que alguém arrastou sem querer".
"""

from datetime import date
from decimal import Decimal

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from cat.dominio.cat42.correcao import CORRECOES_POR_PEDIDO
from cat.infraestrutura.analitico.razao import ESQUEMA_FICHA3
from cat.infraestrutura.planilhas.correcoes_da_planilha import (
    COLUNA_DO_MOTIVO,
    EDITAVEIS,
    PlanilhaIlegivel,
    _sem_acento,
    conferir,
)
from cat.infraestrutura.planilhas.razao import COLUNAS_FICHA3, gerar_ficha3

CHAVE_DA_ENTRADA = "1" * 44
CHAVE_DA_VENDA = "2" * 44
# a planilha tem três linhas de cabeçalho: faixa, título e número do campo
PRIMEIRA_LINHA = 4


def d(v) -> Decimal:
    return Decimal(str(v))


def _linha(**campos) -> dict:
    base = {
        "cnpj": "11111111000191", "codigo": "4002", "descricao": "Xampu 350ml",
        "ncm": "33051000", "unidade_estoque": "UN", "numero": 1, "data": date(2022, 8, 5),
        "especie": "entrada", "devolucao": False, "cfop": "1403", "cst_icms": "060",
        "documento": CHAVE_DA_ENTRADA, "origem": "efd", "enquadramento": None,
        "enquadramento_indefinido": False, "ficha_retirada": False, "corrigida": False,
        "unidade_origem": "UN", "fator_conversao": d(1), "unidade_sem_fator": False,
        "quantidade": d(10), "valor_item": d(100), "icms_suportado": d(20),
        "valor_unitario_usado": d(2), "icms_efetivo": None, "aliquota": None,
        "aliquota_documento": None, "reducao_base": None, "saldo_quantidade": d(10),
        "saldo_unitario": d(2), "saldo_valor": d(20), "ressarcimento": d(0), "complemento": d(0),
        "chave": CHAVE_DA_ENTRADA, "numero_item": 1, "modelo": "55", "participante": "F1",
        "numero_documento": "1", "serie": "1", "codigo_original": None,
        "credito_operacao_propria": d(0),
    }
    base.update(campos)
    return base


@pytest.fixture
def ficha3(tmp_path):
    """Uma entrada e duas vendas a consumidor da mesma mercadoria."""
    linhas = [
        _linha(),
        _linha(numero=2, data=date(2022, 8, 6), especie="saida", cfop="5405", cst_icms="560",
               enquadramento=1, quantidade=d(-2), valor_item=d(30), icms_suportado=d(-4),
               icms_efetivo=d("2.88"), aliquota=d(18), aliquota_documento=d(18), reducao_base=d(0),
               documento=CHAVE_DA_VENDA, chave=CHAVE_DA_VENDA, numero_item=1,
               numero_documento="900", saldo_quantidade=d(8), saldo_valor=d(16)),
        _linha(numero=3, data=date(2022, 8, 7), especie="saida", cfop="5405", cst_icms="560",
               enquadramento=1, quantidade=d(-3), valor_item=d(45), icms_suportado=d(-6),
               icms_efetivo=d("4.32"), aliquota=d(18), aliquota_documento=d(18), reducao_base=d(0),
               documento=CHAVE_DA_VENDA, chave=CHAVE_DA_VENDA, numero_item=2,
               numero_documento="900", saldo_quantidade=d(5), saldo_valor=d(10)),
    ]
    caminho = tmp_path / "ficha3.parquet"
    pq.write_table(pa.Table.from_pydict(
        {c: [l[c] for l in linhas] for c in ESQUEMA_FICHA3.names}, schema=ESQUEMA_FICHA3),
        str(caminho))
    return caminho


@pytest.fixture
def planilha(ficha3, tmp_path):
    destino = tmp_path / "ficha3.xlsx"
    gerar_ficha3(str(ficha3), str(destino))
    return destino


def _coluna(aba, campo: str) -> int:
    """A coluna da planilha pelo título, 1-based, como o openpyxl endereça."""
    titulo = next(c.titulo for c in COLUNAS_FICHA3 if c.campo == campo)
    for celula in aba[2]:
        if celula.value and _sem_acento(celula.value) == _sem_acento(titulo):
            return celula.column
    raise AssertionError(f"coluna {campo} não está na planilha")


def editar(caminho, mudancas: list[tuple[int, str, object]]):
    """Mexe nas células como uma pessoa mexeria: (linha, coluna, valor)."""
    livro = openpyxl.load_workbook(caminho)
    aba = livro.worksheets[0]
    for linha, campo, valor in mudancas:
        aba.cell(row=linha, column=_coluna(aba, campo)).value = valor
    livro.save(caminho)
    return caminho


class TestOQueMudou:
    def test_a_planilha_intocada_nao_propoe_nada(self, planilha, ficha3, tmp_path):
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert r.linhas_lidas == 3
        assert r.correcoes == []
        assert r.erros == []

    def test_aliquota_da_mercadoria_com_o_antes_e_o_depois(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA + 1, "aliquota", 25),
            (PRIMEIRA_LINHA + 1, COLUNA_DO_MOTIVO, "NCM 3305.10.00, art. 55, IV do RICMS"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert len(r.correcoes) == 1
        c = r.correcoes[0]
        assert (c["campo"], c["alvo"], c["codigo"]) == ("aliquota", "mercadoria", "4002")
        assert (c["de"], c["para"]) == ("18.0000", "25.0000")
        assert c["frase"] == "Alíquota interna (mercadoria 4002): 18.0000 → 25.0000"
        assert c["motivo"] == "NCM 3305.10.00, art. 55, IV do RICMS"
        assert c["linha_na_planilha"] == PRIMEIRA_LINHA + 1
        # a correção é da mercadoria: alcança as três linhas dela, não só a editada
        assert c["linhas_atingidas"] == 3

    def test_a_linha_e_achada_pela_chave_e_pelo_item(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA + 2, "cod_legal", 4),
            (PRIMEIRA_LINHA + 2, COLUNA_DO_MOTIVO, "venda interestadual confirmada pelo cliente"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert [(c["campo"], c["de"], c["para"], c["numero_item"]) for c in r.correcoes] == [
            ("enquadramento", "1", "4", 2)]
        assert r.correcoes[0]["onde"] == "documento …222222, item 2"
        assert r.correcoes[0]["linhas_atingidas"] == 1

    def test_tirar_da_ficha_e_uma_correcao_de_linha(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA + 1, "tirar_da_ficha", "sim"),
            (PRIMEIRA_LINHA + 1, COLUNA_DO_MOTIVO, "nota cancelada na SEFAZ fora do prazo"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert [(c["campo"], c["de"], c["para"]) for c in r.correcoes] == [("excluida", "nao", "sim")]

    def test_quantidade_e_valor_da_mesma_linha_viram_duas_correcoes(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA, "qtd_entrada", 12),
            (PRIMEIRA_LINHA, "valor_item", 120),
            (PRIMEIRA_LINHA, "suportado_entrada", 24),
            (PRIMEIRA_LINHA, COLUNA_DO_MOTIVO, "ERP escriturou 10 e a nota traz 12"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert sorted(c["campo"] for c in r.correcoes) == [
            "icms_suportado", "quantidade", "valor_item"]
        quantidade = next(c for c in r.correcoes if c["campo"] == "quantidade")
        assert (quantidade["de"], quantidade["para"]) == ("10.000000", "12.000000")

    def test_a_virgula_decimal_do_excel_e_entendida(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA + 1, "reducao_base", "33,33"),
            (PRIMEIRA_LINHA + 1, COLUNA_DO_MOTIVO, "carga de 12% do art. 34 do Anexo II"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert [(c["campo"], c["para"]) for c in r.correcoes] == [("reducao_base", "33.3300")]

    def test_diferenca_de_arredondamento_nao_e_correcao(self, planilha, ficha3, tmp_path):
        # o Excel devolve 18 como 18.000000001: isso é o float dele, não correção
        editar(planilha, [
            (PRIMEIRA_LINHA + 1, "aliquota", 18.000000001),
            (PRIMEIRA_LINHA + 1, COLUNA_DO_MOTIVO, "conferido e mantido"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert r.correcoes == []
        assert "nenhum valor mudou" in r.avisos[0]["mensagem"]


class TestOQueEleRecusa:
    def test_mudanca_sem_motivo_volta_como_erro(self, planilha, ficha3, tmp_path):
        editar(planilha, [(PRIMEIRA_LINHA + 1, "aliquota", 25)])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert r.correcoes == []
        assert r.erros[0]["linha"] == PRIMEIRA_LINHA + 1
        assert r.erros[0]["rotulo"] == "Alíquota interna"
        assert "motivo escrito" in r.erros[0]["mensagem"]

    def test_valor_que_o_dominio_recusa_volta_com_a_linha(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA + 1, "cod_legal", 9),
            (PRIMEIRA_LINHA + 1, COLUNA_DO_MOTIVO, "enquadramento revisto com o cliente"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert r.correcoes == []
        assert "de 0 a 4" in r.erros[0]["mensagem"]

    def test_a_mesma_mercadoria_com_dois_valores_nao_grava_nenhum(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA + 1, "aliquota", 25),
            (PRIMEIRA_LINHA + 1, COLUNA_DO_MOTIVO, "art. 55, IV"),
            (PRIMEIRA_LINHA + 2, "aliquota", 12),
            (PRIMEIRA_LINHA + 2, COLUNA_DO_MOTIVO, "art. 34 do Anexo II"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert r.correcoes == []
        assert "Deixe um valor só" in r.erros[0]["mensagem"]

    def test_linha_que_nao_existe_na_ficha_volta_como_erro(self, planilha, ficha3, tmp_path):
        editar(planilha, [
            (PRIMEIRA_LINHA + 1, "chave", "9" * 44),
            (PRIMEIRA_LINHA + 1, "aliquota", 25),
            (PRIMEIRA_LINHA + 1, COLUNA_DO_MOTIVO, "alíquota do art. 55, IV"),
        ])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert r.correcoes == []
        assert r.nao_achadas == 1
        assert "Não achei esta linha" in r.erros[0]["mensagem"]

    def test_linha_sem_motivo_que_nao_casa_nao_vira_erro(self, planilha, ficha3, tmp_path):
        """Planilha de outro razão tem milhões de linhas que não casam: acusar
        cada uma seria dar à pessoa um relatório que ela não consegue ler."""
        editar(planilha, [(PRIMEIRA_LINHA + 1, "chave", "9" * 44)])
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert (r.erros, r.nao_achadas) == ([], 0)

    def test_planilha_de_outra_coisa_e_recusada(self, ficha3, tmp_path):
        outra = tmp_path / "outra.xlsx"
        livro = openpyxl.Workbook()
        livro.active.append(["Produto", "Quantidade", "Valor"])
        livro.active.append(["Xampu", 1, 2])
        livro.save(outra)
        with pytest.raises(PlanilhaIlegivel, match="cabeçalho da Ficha 3"):
            conferir(str(outra), str(ficha3), pasta=str(tmp_path / "rascunho"))

    def test_extensao_que_nao_e_planilha_e_recusada(self, ficha3, tmp_path):
        with pytest.raises(PlanilhaIlegivel, match="xlsx ou csv"):
            conferir(str(tmp_path / "ficha3.pdf"), str(ficha3), pasta=str(tmp_path / "rascunho"))


class TestOCsvTambemServe:
    def test_o_csv_gerado_volta_e_e_entendido(self, ficha3, tmp_path):
        csv_gerado = tmp_path / "ficha3.csv"
        gerar_ficha3(str(ficha3), str(csv_gerado), formato="csv")
        texto = csv_gerado.read_text(encoding="utf-8-sig")
        linhas = texto.splitlines()
        colunas = linhas[0].split(";")
        # o CSV junta título e número do campo num cabeçalho só
        alvo = colunas.index("Alíquota do Confronto (%)")
        motivo = colunas.index("Motivo da Correção")
        celulas = linhas[2].split(";")
        celulas[alvo] = "25,00"
        celulas[motivo] = "NCM 3305.10.00, art. 55, IV"
        linhas[2] = ";".join(celulas)
        csv_gerado.write_text("\n".join(linhas), encoding="utf-8-sig")

        r = conferir(str(csv_gerado), str(ficha3), pasta=str(tmp_path / "rascunho"))

        assert [(c["campo"], c["de"], c["para"]) for c in r.correcoes] == [
            ("aliquota", "18.0000", "25.0000")]


class TestOContrato:
    def test_toda_coluna_editavel_existe_no_leiaute(self):
        do_leiaute = {c.campo for c in COLUNAS_FICHA3}
        assert {e.coluna for e in EDITAVEIS} <= do_leiaute
        assert COLUNA_DO_MOTIVO in do_leiaute

    def test_o_limite_do_pedido_e_o_do_dominio(self, planilha, ficha3, tmp_path):
        r = conferir(str(planilha), str(ficha3), pasta=str(tmp_path / "rascunho"))
        assert r.como_json()["limite"] == CORRECOES_POR_PEDIDO
