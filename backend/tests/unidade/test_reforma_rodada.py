"""A rodada da simulação: das planilhas do lote às visões, com exclusão,
pendência e o reconhecimento das planilhas na importação."""

from datetime import datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from cat.dominio.lote import TipoDeArquivo
from cat.infraestrutura.analitico.reforma import serializar, simular
from cat.infraestrutura.arquivos.classificador import classificar
from cat.infraestrutura.planilhas.reforma import gerar

CABECALHO_DA_BASE = [
    "NFCC_LOJ", "cnpj", "Estado ", "dt_ini", "cod_item", "descr_item", "qtd_c", "vl_item",
    "Valor Líquido", "Status Custo", "Base da Análise", "Receita Bruta", "Débito ICMS",
    "Débito IPI", "Débito PIS/COFINS", "Receita Líquida",
    "Custo Líquido (sem ICMS, IPI, PIS/COFINS)", "IPI na Entrada (custo - não recuperável)",
    "PIS/COFINS sem Crédito CST 70/72/73/75 (custo)", "CMV", "Crédito ICMS (entrada)",
    "Crédito PIS/COFINS (entrada)",
]


def _venda(loja, uf, item, descr, rb, icms, pc, custo, credito_icms, status="Loja + Item",
           pcnc=0, ipi=0, credito_pc=0):
    rl = rb - icms - pc
    sem_custo = custo is None
    return [loja, "44.000.002/0001-37", uf, datetime(2025, 12, 1), item, descr, 1, rb, rl,
            status, "Sem custo - fora da análise" if sem_custo else status,
            None if sem_custo else rb, None if sem_custo else icms, None if sem_custo else 0,
            None if sem_custo else pc, None if sem_custo else rl, custo,
            None if sem_custo else ipi, None if sem_custo else pcnc, None,
            None if sem_custo else credito_icms, None if sem_custo else credito_pc]


@pytest.fixture
def pasta(tmp_path):
    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Saídas Dez."
    aba.append(CABECALHO_DA_BASE)
    aba.append(_venda(1, "BA", 10, "CERV LT 350ML", 100, 18, 9.25, 60, 10, credito_pc=4))
    aba.append(_venda(2, "BA", 10, "CERV LT 350ML", 50, 9, 4.62, 30, 5))
    aba.append(_venda(1, "SE", 20, "ARROZ 5KG", 200, 14, 0, 150, 12, pcnc=3))
    aba.append(_venda(1, "BA", 30, "OLEO SOJA 900ML", 80, 9.6, 1, 60, 5))
    aba.append(_venda(1, "BA", 40, "PAO FRANCES KG", 40, 0, 0, None, 0,
                      status="Sem compra Abr-Dez"))
    aba.append(_venda(1, "BA", 50, "ITEM SEM CADASTRO", 30, 5, 0, 20, 2))
    aba.append([None] * len(CABECALHO_DA_BASE))
    livro.save(tmp_path / "saidas.xlsx")

    cadastro = openpyxl.Workbook()
    aba = cadastro.active
    aba.append(["Empresa: empresa 12"])
    aba.append([])
    aba.append([None, "TIPO", "CODACESSO", "SEQPRODUTO", "DESCCOMPLETA", "NCM Original",
                "NCM BMS", "Validação", "CST", "cClassTrib", "Redução de Alíquota", "Base Legal"])
    aba.append([None, "EAN", "1", 10, "CERV LT 350ML", "22030000", "22030000", None, "000",
                "000001", "Não aplicável", ""])
    aba.append([None, "EAN", "2", 20, "ARROZ 5KG", "10063021", "10063021", None, 200,
                200003, 1, "Anexo I"])
    aba.append([None, "EAN", "3", 30.0, "OLEO SOJA 900ML", "15079011", "15079011", None, 200,
                200034, 0.6, "Anexo VII"])
    cadastro.save(tmp_path / "classificacao.xlsx")
    return tmp_path


class TestReconhecimento:
    def test_a_base_pelo_cabecalho(self, pasta):
        a = classificar(str(pasta / "saidas.xlsx"))
        assert a.tipo is TipoDeArquivo.REFORMA_BASE_CUSTEADA
        assert a.cnpj == "44000002000137"
        assert a.competencia.isoformat() == "2025-12-01"

    def test_a_classificacao_com_titulo_antes(self, pasta):
        assert classificar(str(pasta / "classificacao.xlsx")).tipo is \
            TipoDeArquivo.REFORMA_CLASSIFICACAO

    def test_planilha_quase_diz_o_que_falta(self, tmp_path):
        livro = openpyxl.Workbook()
        livro.active.append([c for c in CABECALHO_DA_BASE if c != "Débito ICMS"])
        livro.save(tmp_path / "quase.xlsx")
        a = classificar(str(tmp_path / "quase.xlsx"))
        assert a.tipo is TipoDeArquivo.DESCONHECIDO
        assert "debito_icms" in a.motivo

    def test_resultado_sem_cst_nao_vira_cadastro(self, tmp_path):
        """A aba "Base por SKU" da planilha de resultado tem item, NCM e
        cClassTrib, mas não CST: lida como cadastro, entraria em dobro."""
        livro = openpyxl.Workbook()
        livro.active.append(["SEQPRODUTO", "Descrição", "NCM", "cClassTrib", "Grupo"])
        livro.save(tmp_path / "resultado.xlsx")
        assert classificar(str(tmp_path / "resultado.xlsx")).tipo is TipoDeArquivo.DESCONHECIDO

    def test_a_reforma_alimenta_cbs_e_ibs(self):
        assert TipoDeArquivo.REFORMA_BASE_CUSTEADA.alimenta("cbs")
        assert TipoDeArquivo.REFORMA_CLASSIFICACAO.alimenta("ibs")
        assert not TipoDeArquivo.REFORMA_BASE_CUSTEADA.alimenta("icms")


class TestSimulacao:
    @pytest.fixture
    def resultado(self, pasta):
        return simular([str(pasta / "saidas.xlsx")], [str(pasta / "classificacao.xlsx")])

    def test_soma_por_competencia_uf_e_item(self, resultado):
        """As duas lojas da BA que venderam o item 10 viram uma célula."""
        assert resultado.linhas_lidas == 6
        assert resultado.linhas_consideradas == 4
        assert resultado.total.base.receita_bruta == Decimal(430)
        assert len(resultado.enquadramentos) == 3

    def test_sem_custo_e_exclusao_pelo_motivo_da_base(self, resultado):
        assert set(resultado.exclusoes) == {"Sem compra Abr-Dez"}
        assert resultado.exclusoes["Sem compra Abr-Dez"].receita_liquida == Decimal(40)

    def test_sem_cadastro_e_pendencia_e_nao_aliquota_padrao(self, resultado):
        pendencia = resultado.pendencias["Item sem cadastro de classificação"]
        assert pendencia.linhas == 1 and pendencia.receita_bruta == Decimal(30)
        assert "50" in pendencia.itens

    def test_grupos_e_seletivo(self, resultado):
        grupos = {l.chave: l for l in resultado.visoes["grupo"]}
        assert set(grupos) == {"padrao", "zero", "reducao_60"}
        # a cerveja é bebida alcoólica: IS = custo × 50%
        assert grupos["padrao"].base.seletivo_pleno == Decimal(45)
        assert grupos["zero"].base.seletivo_pleno == 0

    def test_2025_e_o_modelo_de_hoje(self, resultado):
        t = resultado.total
        assert t.anos[2025].lucro_bruto == t.base.lucro_bruto_atual

    def test_visoes_somam_o_total(self, resultado):
        """Somar as chaves de qualquer visão dá o total — o modelo é linear."""
        total = resultado.total.anos[2033].lucro_bruto
        for visao in ("grupo", "uf", "capitulo", "posicao", "familia", "sku",
                      "categoria_seletivo", "competencia"):
            soma = sum(l.anos[2033].lucro_bruto for l in resultado.visoes[visao])
            assert soma == pytest.approx(total, abs=Decimal("0.000001")), visao

    def test_resumo_e_planilha(self, resultado, tmp_path):
        resumo = serializar(resultado)
        assert [a["ano"] for a in resumo["por_ano"]] == list(range(2025, 2034))
        assert resumo["base"]["competencias"] == ["2025-12"]
        assert resumo["pendencias"][0]["skus"] == 1
        destino = tmp_path / "simulacao.xlsx"
        gerar(resultado, str(destino))
        abas = openpyxl.load_workbook(destino, read_only=True).sheetnames
        assert abas[0] == "Premissas" and "DRE 2025-2033" in abas and "Fora da conta" in abas

    def test_cancelamento_e_respeitado(self, pasta):
        from cat.infraestrutura.analitico.escrita import LeituraCancelada

        def parar():
            raise LeituraCancelada()
        with pytest.raises(LeituraCancelada):
            simular([str(pasta / "saidas.xlsx")], [str(pasta / "classificacao.xlsx")],
                    deve_parar=parar)


class TestPreparar:
    @pytest.fixture
    def sessao(self):
        from cat.infraestrutura.repositorios.modelos import Base
        motor = create_engine("sqlite://")
        Base.metadata.create_all(motor)
        with Session(motor) as s:
            yield s

    def _projeto(self, sessao, *tipos):
        from cat.infraestrutura.repositorios.modelos import (
            ArquivoDoLoteDB, EmpresaDB, LoteDB, ProjetoDB)
        from datetime import date
        empresa = EmpresaDB(razao_social="EMPRESA DO TESTE LTDA", cnpj_raiz="44000002")
        sessao.add(empresa)
        sessao.flush()
        projeto = ProjetoDB(empresa_id=empresa.id, frente="reforma", nome="Reforma",
                            competencia_ini=date(2025, 12, 1),
                            competencia_fim=date(2025, 12, 31))
        sessao.add(projeto)
        sessao.flush()
        lote = LoteDB(projeto_id=projeto.id, pasta="C:/x")
        sessao.add(lote)
        sessao.flush()
        for i, tipo in enumerate(tipos):
            sessao.add(ArquivoDoLoteDB(lote_id=lote.id, caminho=f"C:/x/{i}.xlsx",
                                       nome=f"{i}.xlsx", tamanho=1, tipo=tipo.value))
        sessao.commit()
        return projeto.id

    def test_sem_classificacao_recusa_e_diz_o_que_falta(self, sessao):
        from cat.aplicacao.casos_de_uso.simular_reforma import NadaParaSimular, preparar
        projeto = self._projeto(sessao, TipoDeArquivo.REFORMA_BASE_CUSTEADA)
        with pytest.raises(NadaParaSimular, match="classificação"):
            preparar(projeto, 1, sessao)

    def test_com_as_duas_entra_na_fila(self, sessao):
        from cat.aplicacao.casos_de_uso.simular_reforma import ETAPA, preparar
        projeto = self._projeto(sessao, TipoDeArquivo.REFORMA_BASE_CUSTEADA,
                                TipoDeArquivo.REFORMA_CLASSIFICACAO)
        execucao = preparar(projeto, 1, sessao)
        assert execucao.etapa == ETAPA and execucao.situacao == "na_fila"
        assert execucao.arquivos_totais == 2
