"""A Ficha 3 no leiaute do papel de trabalho, e o cabeçalho padrão do projeto.

O leiaute é o que o cliente confere contra a entrega de outro escritório: faixa
por bloco, título e o número do campo da CAT 42 embaixo. O cabeçalho azul com
Arial 10 branco vale para **todas** as planilhas do projeto (pedido do Victor,
18/09/2026), e por isso o teste do estilo olha uma planilha comum, não só esta.
"""

import csv
from datetime import date
from decimal import Decimal

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from cat.infraestrutura.analitico.razao import ESQUEMA_FICHA3, ESQUEMA_FICHAS
from cat.infraestrutura.planilhas.conferencia import (
    ALTURA_DO_CABECALHO,
    AZUL_DO_CABECALHO,
    Coluna,
    _faixas,
)
from cat.infraestrutura.planilhas.razao import COLUNAS_FICHA3, gerar_ficha3, gerar_fichas


def d(v) -> Decimal:
    return Decimal(str(v))


def _linha(**campos) -> dict:
    base = {
        "cnpj": "11111111000191", "codigo": "X", "descricao": "Refrigerante cola 2L",
        "ncm": "22021000", "unidade_estoque": "UN", "numero": 1, "data": date(2021, 1, 5),
        "especie": "entrada", "devolucao": False, "cfop": "1403", "cst_icms": "060",
        "documento": "4" * 44, "origem": "efd", "enquadramento": None,
        "enquadramento_indefinido": False, "ficha_retirada": False, "unidade_origem": "UN",
        "fator_conversao": d(1), "unidade_sem_fator": False, "quantidade": d(10),
        "valor_item": d(100), "icms_suportado": d(20), "valor_unitario_usado": d(2),
        "icms_efetivo": None, "aliquota": None, "reducao_base": None,
        "saldo_quantidade": d(10), "saldo_unitario": d(2), "saldo_valor": d(20),
        "ressarcimento": d(0), "complemento": d(0), "chave": "4" * 44, "numero_item": 1,
        "modelo": "55", "participante": "F1", "numero_documento": "1", "serie": "1",
        "codigo_original": None, "credito_operacao_propria": d(0),
    }
    base.update(campos)
    return base


@pytest.fixture
def ficha3(tmp_path):
    """Uma entrada, uma venda a consumidor e uma saída para outro estado."""
    linhas = [
        _linha(),
        _linha(numero=2, data=date(2021, 1, 6), especie="saida", cfop="5405", cst_icms="560",
               enquadramento=1, quantidade=d(-2), valor_item=d(30), icms_suportado=d(-4),
               icms_efetivo=d("2.88"), aliquota=d(18), reducao_base=d(52),
               saldo_quantidade=d(8), saldo_valor=d(16), complemento=d(0), ressarcimento=d("1.12")),
        _linha(numero=3, data=date(2021, 2, 8), especie="saida", cfop="6102", cst_icms="500",
               enquadramento=4, quantidade=d(-1), valor_item=d(9), icms_suportado=d(-2),
               icms_efetivo=d("0.63"), saldo_quantidade=d(-1), saldo_valor=d(14),
               ressarcimento=d("1.37"), credito_operacao_propria=d("0.63")),
    ]
    caminho = tmp_path / "ficha3.parquet"
    pq.write_table(pa.Table.from_pydict(
        {c: [l[c] for l in linhas] for c in ESQUEMA_FICHA3.names}, schema=ESQUEMA_FICHA3),
        str(caminho))
    return caminho


@pytest.fixture
def planilha(ficha3, tmp_path):
    caminho = tmp_path / "ficha3.xlsx"
    assert gerar_ficha3(str(ficha3), str(caminho)) == 3
    return openpyxl.load_workbook(str(caminho)).active


def _coluna(aba, titulo: str) -> int:
    titulos = [c.value for c in aba[2]]
    return titulos.index(titulo) + 1


class TestCabecalhoEmBlocos:
    def test_tres_linhas_faixa_titulo_e_numero(self, planilha):
        assert planilha.cell(row=1, column=1).value == "Dados Gerais"
        assert planilha.cell(row=2, column=1).value == "Período"
        assert planilha.cell(row=3, column=_coluna(planilha, "Data")).value == "(2)"
        assert planilha.cell(row=3, column=_coluna(planilha, "Chave")).value == "(3)"
        # o dado começa depois das três linhas, e o painel congela ali
        assert planilha.cell(row=4, column=_coluna(planilha, "Código da Mercadoria")).value == "X"
        assert planilha.freeze_panes == "A4"

    def test_a_faixa_cobre_o_bloco_inteiro(self, planilha):
        mescladas = {str(m) for m in planilha.merged_cells.ranges}
        largura = len([c for c in COLUNAS_FICHA3 if c.bloco == "Dados Gerais"])
        assert f"A1:{openpyxl.utils.get_column_letter(largura)}1" in mescladas
        blocos = {planilha.cell(row=1, column=i).value for i in range(1, len(COLUNAS_FICHA3) + 1)}
        assert blocos >= {"Dados Gerais", "Entradas", "Saídas", "Valor de Confronto",
                          "Saldo", "Apuração", "Inconsistências"}

    def test_faixas_junta_colunas_seguidas_do_mesmo_bloco(self):
        colunas = (Coluna("a", "A", bloco="Um"), Coluna("b", "B", bloco="Um"),
                   Coluna("c", "C", bloco="Dois"))
        assert _faixas(colunas) == [("Um", 0, 1), ("Dois", 2, 2)]

    def test_todo_campo_do_leiaute_esta_numerado_uma_vez(self):
        numeros = [c.numero for c in COLUNAS_FICHA3 if c.numero]
        assert numeros == [str(n) for n in range(1, 28)]


class TestEstiloPadrao:
    def test_o_cabecalho_e_o_mesmo_em_qualquer_planilha(self, ficha3, tmp_path):
        """O estilo é do projeto, não da Ficha 3: a planilha de fichas usa o mesmo."""
        vazio = tmp_path / "fichas.parquet"
        pq.write_table(pa.Table.from_pydict(
            {c: [] for c in ESQUEMA_FICHAS.names}, schema=ESQUEMA_FICHAS), str(vazio))
        caminho = tmp_path / "fichas.xlsx"
        gerar_fichas(str(vazio), str(caminho))
        aba = openpyxl.load_workbook(str(caminho)).active
        celula = aba.cell(row=1, column=1)
        assert celula.value == "CNPJ do estabelecimento"      # uma linha só, como era
        assert aba.freeze_panes == "A2"
        assert celula.fill.start_color.rgb == "FF" + AZUL_DO_CABECALHO.lstrip("#")
        assert (celula.font.name, celula.font.sz, celula.font.b) == ("Arial", 10.0, True)
        assert celula.font.color.rgb == "FFFFFFFF"
        assert (celula.alignment.horizontal, celula.alignment.vertical) == ("center", "center")
        assert celula.alignment.wrap_text

    def test_a_faixa_de_titulos_tem_altura_para_tres_linhas(self, planilha):
        assert planilha.row_dimensions[2].height == ALTURA_DO_CABECALHO


class TestOQueCaiEmCadaColuna:
    def test_dados_gerais(self, planilha):
        periodo = planilha.cell(row=4, column=_coluna(planilha, "Período")).value
        assert periodo.date() == date(2021, 1, 1)      # o mês da linha, dia 1
        assert planilha.cell(row=4, column=_coluna(planilha, "NCM")).value == "22021000"
        assert planilha.cell(row=4, column=_coluna(planilha, "Descrição")).value == "Refrigerante cola 2L"
        assert planilha.cell(row=4, column=_coluna(planilha, "VL_ITEM (base de cálculo do ICMS)")).value == 100
        # sem código de origem, o "Código Original" repete o da mercadoria
        assert planilha.cell(row=4, column=_coluna(planilha, "Código Original")).value == "X"

    def test_tipo_0_entrada_1_saida(self, planilha):
        coluna = _coluna(planilha, "Tipo 0 - Entrada 1 - Saída")
        assert [planilha.cell(row=l, column=coluna).value for l in (4, 5, 6)] == ["0", "1", "1"]

    def test_entrada_e_saida_em_colunas_separadas(self, planilha):
        entrada = _coluna(planilha, "Quantidade (Entrada)")
        saida = _coluna(planilha, "Quantidade (Saída)")
        assert planilha.cell(row=4, column=entrada).value == 10
        assert planilha.cell(row=4, column=saida).value is None
        assert planilha.cell(row=5, column=entrada).value is None
        assert planilha.cell(row=5, column=saida).value == 2      # sem o sinal da baixa

    def test_o_valor_da_saida_cai_na_coluna_do_enquadramento(self, planilha):
        enq1 = _coluna(planilha, "Saída a Consumidor ou Usuário Final - Código Enquadramento 1")
        enq4 = _coluna(planilha, "Saída para Outro Estado - Código Enquadramento 4")
        assert (planilha.cell(row=5, column=enq1).value,
                planilha.cell(row=5, column=enq4).value) == (30, None)
        assert (planilha.cell(row=6, column=enq1).value,
                planilha.cell(row=6, column=enq4).value) == (None, 9)

    def test_o_confronto_separa_saida_de_entrada(self, planilha):
        saida = _coluna(planilha, "ICMS Efetivo na Saída a Consumidor ou Usuário Final")
        entrada = _coluna(planilha, "ICMS Efetivo da Entrada nas Demais Hipóteses")
        assert planilha.cell(row=5, column=saida).value == 2.88     # enquadramento 1
        assert planilha.cell(row=5, column=entrada).value is None
        assert planilha.cell(row=6, column=entrada).value == 0.63   # enquadramento 4
        assert planilha.cell(row=6, column=saida).value is None

    def test_as_nossas_colunas_de_confronto(self, planilha):
        assert planilha.cell(row=5, column=_coluna(planilha, "Alíquota do ICMS")).value == 18
        assert planilha.cell(row=5, column=_coluna(planilha, "Redução de Base (%)")).value == 52
        assert planilha.cell(row=5, column=_coluna(planilha, "CST")).value == "560"

    def test_inconsistencias_marcadas_na_linha(self, planilha):
        negativo = _coluna(planilha, "Saldo Parcial Negativo")
        assert planilha.cell(row=6, column=negativo).value == "Sim"
        assert planilha.cell(row=5, column=negativo).value == "Não"


class TestCompatibilidade:
    def test_ficha_sem_as_colunas_novas_sai_com_elas_vazias(self, ficha3, tmp_path):
        """Rodada antiga não tem NCM nem VL_ITEM: a planilha sai, com a coluna vazia."""
        antiga = tmp_path / "antiga.parquet"
        tabela = pq.read_table(str(ficha3)).drop(["ncm", "valor_item", "descricao"])
        pq.write_table(tabela, str(antiga))
        caminho = tmp_path / "antiga.xlsx"
        assert gerar_ficha3(str(antiga), str(caminho)) == 3
        aba = openpyxl.load_workbook(str(caminho)).active
        assert aba.cell(row=4, column=_coluna(aba, "NCM")).value is None
        assert aba.cell(row=4, column=_coluna(aba, "Código da Mercadoria")).value == "X"

    def test_csv_leva_o_numero_junto_do_titulo(self, ficha3, tmp_path):
        caminho = tmp_path / "ficha3.csv"
        assert gerar_ficha3(str(ficha3), str(caminho), formato="csv") == 3
        with open(caminho, encoding="utf-8-sig", newline="") as f:
            cabecalho = next(csv.reader(f, delimiter=";"))
        assert "Chave (3)" in cabecalho and "Valor do Ressarcimento (25)" in cabecalho
        assert "Período" in cabecalho          # sem número, fica só o título
