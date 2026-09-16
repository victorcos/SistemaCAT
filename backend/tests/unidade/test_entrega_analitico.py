"""A montagem da entrega sobre o cenário da etapa 7.

Loja de SP: janeiro apto e limpo vai para o envio; fevereiro sai como prévia.
Loja do PR: fora de SP. A entrega junta os três no relatório e leva só janeiro
para o dossiê.
"""

import csv
import io
import os
import zipfile
from datetime import datetime, timezone

import openpyxl
import pytest

from cat.infraestrutura.analitico.arquivo_digital import PASTA_ENVIO, gerar
from cat.infraestrutura.analitico.entrega import (
    ARQUIVO_MANIFESTO,
    ARQUIVO_PACOTE,
    ARQUIVO_RELATORIO,
    PASTA_DOSSIE,
    ArquivoAlterado,
    Contexto,
    Fontes,
    estabelecimentos,
    montar,
    serializar,
)
from cat.infraestrutura.analitico.suportado import ApuracaoCancelada
from tests.unidade.test_arquivo_digital_analitico import PR, SP, fontes  # noqa: F401  (fixture)


@pytest.fixture
def etapa7(fontes, tmp_path):  # noqa: F811
    pasta = tmp_path / "arquivo_digital"
    pasta.mkdir()
    r = gerar(fontes, str(pasta))
    assert (r.para_envio, r.previas) == (1, 1)
    return Fontes(arquivo_digital=str(pasta), apuracao=fontes.apuracao, razao=fontes.razao)


def contexto() -> Contexto:
    return Contexto(
        empresa="LOJA DE TESTE", cnpj_matriz=SP, trabalho="CAT 42 2024", periodo="01/2024 a 02/2024",
        venda_a_consumidor="Enquadramento 1", gerado_por="Revisora", versao_do_sistema="9.9.9",
        gerado_em=datetime(2026, 9, 16, 15, 0, tzinfo=timezone.utc),
        resumos={"arquivo_digital": {"entradas_sem_icms": 3, "por_trava": []},
                 "apuracao": {"por_motivo": [{"codigo": "sem_inventario", "rotulo": "Sem inventário",
                                              "o_que_fazer": "Importar o bloco H.", "competencias": 1}]}},
        trilha=[{"etapa": "arquivo_digital", "execucao_id": 51, "iniciada_por": "Analista",
                 "terminada_em": datetime(2026, 9, 16, 14, 0, tzinfo=timezone.utc), "segundos": 435.6}],
    )


@pytest.fixture
def montada(etapa7, tmp_path):
    destino = tmp_path / "entrega"
    return destino, montar(etapa7, str(destino), contexto())


class TestOResumo:
    def test_conta_cada_competencia_onde_ficou(self, montada):
        _, r = montada
        s = serializar(r)
        assert (s["competencias"], s["competencias_de_sp"], s["para_envio"], s["previas"], s["fora_de_sp"]) == (3, 2, 1, 1, 1)
        assert (s["estabelecimentos"], s["estabelecimentos_no_dossie"]) == (2, 1)
        assert s["ressarcimento_para_envio"] == "1.50"
        # o apurado soma as três, inclusive a do PR
        assert s["ressarcimento"] == "10.50"

    def test_pendencias_vem_dos_resumos(self, montada):
        _, r = montada
        s = serializar(r)
        assert [p["codigo"] for p in s["pendencias"]] == ["sem_inventario", "entradas_sem_icms"]
        assert s["por_gravidade"] == {"trava": 1, "atencao": 1, "informacao": 0}


class TestODossie:
    def test_so_o_que_vai_a_sefaz(self, montada, etapa7):
        destino, _ = montada
        pasta = destino / PASTA_DOSSIE
        assert sorted(os.listdir(pasta)) == [SP]
        txt = os.listdir(pasta / SP / "arquivo_digital")
        assert txt == [f"CAT5_SP_{SP}_1_2024.txt"]
        original = open(os.path.join(etapa7.arquivo_digital, PASTA_ENVIO, txt[0]), "rb").read()
        assert (pasta / SP / "arquivo_digital" / txt[0]).read_bytes() == original

    def test_ficha3_e_planilhas_so_das_competencias_de_envio(self, montada):
        destino, _ = montada
        pasta = destino / PASTA_DOSSIE / SP
        assert {f"saldos_1050_{SP}.xlsx", f"apuracao_{SP}.xlsx", f"pre_validacao_{SP}.xlsx",
                f"ficha3_{SP}.csv", "arquivo_digital"} == set(os.listdir(pasta))
        linhas = list(csv.reader(io.StringIO((pasta / f"ficha3_{SP}.csv").read_text("utf-8-sig")), delimiter=";"))
        # cabeçalho e as duas linhas de janeiro; a venda de fevereiro não entra
        assert len(linhas) == 3
        saldos = openpyxl.load_workbook(pasta / f"saldos_1050_{SP}.xlsx", read_only=True).active
        assert [r[1] for r in saldos.iter_rows(min_row=2, values_only=True)] == ["2024-01"]

    def test_arquivo_mudado_depois_de_pre_validado_nao_entra(self, etapa7, tmp_path):
        caminho = os.path.join(etapa7.arquivo_digital, PASTA_ENVIO, f"CAT5_SP_{SP}_1_2024.txt")
        with open(caminho, "ab") as f:
            f.write(b"1050|X|1,000|0,00|1,000|0,00\r\n")
        with pytest.raises(ArquivoAlterado):
            montar(etapa7, str(tmp_path / "entrega"), contexto())


class TestORelatorioEOPacote:
    def test_relatorio_tem_todas_as_competencias(self, montada):
        destino, _ = montada
        livro = openpyxl.load_workbook(destino / ARQUIVO_RELATORIO, read_only=True)
        assert livro.sheetnames == ["Resumo", "Filial x competência", "Por filial", "Por mês", "Pendências", "Trilha"]
        linhas = list(livro["Filial x competência"].iter_rows(min_row=2, values_only=True))
        assert [(l[0], l[2], l[3]) for l in linhas] == [
            (SP, "2024-01", "Pronta para envio"),
            (SP, "2024-02", "Prévia (com pendência)"),
            (PR, "2024-01", "Fora de SP"),
        ]
        assert "Sem inventário no mês" in linhas[1][7]

    def test_manifesto_confere_cada_arquivo(self, montada):
        destino, _ = montada
        texto = (destino / ARQUIVO_MANIFESTO).read_text("utf-8")
        assert "execucao #51" in texto and "Sistema CAT 9.9.9" in texto
        listados = [l.strip().split(" | ") for l in texto.splitlines() if l.startswith("  ") and " | " in l]
        assert {c for c, _, _ in listados} == {
            ARQUIVO_RELATORIO, f"dossie/{SP}/arquivo_digital/CAT5_SP_{SP}_1_2024.txt",
            f"dossie/{SP}/ficha3_{SP}.csv", f"dossie/{SP}/saldos_1050_{SP}.xlsx",
            f"dossie/{SP}/apuracao_{SP}.xlsx", f"dossie/{SP}/pre_validacao_{SP}.xlsx"}

    def test_pacote_leva_manifesto_relatorio_e_dossie_e_nunca_previa(self, montada):
        destino, r = montada
        nomes = zipfile.ZipFile(destino / ARQUIVO_PACOTE).namelist()
        assert nomes[:2] == [ARQUIVO_MANIFESTO, ARQUIVO_RELATORIO]
        assert len(nomes) == r.arquivos_no_pacote == 7
        assert not any("PREVIA" in n for n in nomes)
        assert len(r.sha256_do_pacote) == 64 and r.bytes_do_pacote == os.path.getsize(destino / ARQUIVO_PACOTE)


class TestATela:
    def test_estabelecimentos_do_dossie_primeiro(self, montada):
        destino, _ = montada
        lista = estabelecimentos(str(destino))
        assert [(e["cnpj"], e["no_dossie"], e["arquivos_no_dossie"]) for e in lista["linhas"]] == [
            (SP, True, 1), (PR, False, 0)]
        assert estabelecimentos(str(destino), so="fora_de_sp")["total"] == 1
        with pytest.raises(ValueError):
            estabelecimentos(str(destino), so="qualquer")


class TestCancelar:
    def test_nao_deixa_pacote_pela_metade(self, etapa7, tmp_path):
        destino = tmp_path / "cancelada"
        with pytest.raises(ApuracaoCancelada):
            montar(etapa7, str(destino), contexto(), deve_parar=lambda: True)
        assert not (destino / ARQUIVO_PACOTE).exists()
        assert not (destino / PASTA_DOSSIE).exists()
