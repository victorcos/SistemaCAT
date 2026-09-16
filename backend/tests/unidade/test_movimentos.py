"""Extração e consolidação do histórico de movimentação.

A EFD de teste reproduz o que os arquivos reais mostraram: a entrada tem
C170, a NF-e própria de saída e o cupom SAT NÃO têm — vão só com o analítico.
O cadastro 0200 se repete entre períodos, e vale o mais recente. Cada
movimento sai marcado pelo que a conferência achou.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import duckdb
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from cat.dominio.cat42.conferencia import Fatia
from cat.dominio.cat42.movimentacao import ResumoDaMovimentacao
from cat.infraestrutura.analitico.movimentacao import (
    ARQUIVO_ANALITICO,
    ARQUIVO_ITENS,
    ARQUIVO_MOVIMENTOS,
    consolidar,
)
from cat.infraestrutura.analitico.movimentos import (
    ARQUIVO_DOCUMENTOS,
    ARQUIVO_INVENTARIO,
    extrair_movimentos,
)
from cat.infraestrutura.planilhas.movimentacao import (
    gerar_analitico,
    gerar_inventario,
    gerar_itens,
    gerar_movimentos,
)

CNPJ = "11517841000278"
CHAVE_ENTRADA = "41210511517841000278550010000446231411953289"
CHAVE_SAIDA = "41210511517841000278550010000446240000000009"
CHAVE_CUPOM = "35210711517841005407590003535520625376456954"

CABECALHO = (f"|0000|015|0|01052021|31052021|EMPRESA DE TESTE|{CNPJ}||PR"
             "|9030138187|4106902||||")
ITEM_A_V1 = "|0200|1000144|Iog Vidativa 160g|7898194090401||CX|00|04031000||04||18|1702200|"
ITEM_A_V2 = "|0200|1000144|Iog Vidativa 160g Ameixa|7898194090401||CX|00|04031000||04||18|1702200|"
ITEM_B = "|0200|537861|Abobora P/Doce Kg Frac|||KG|00|07099990||04||18||"
CONV_A_FD = "|0220|FD|12|"
CONV_A_FD_ANTIGO = "|0220|FD|6|"

C100_ENTRADA = (f"|C100|0|1|F001|55|00|001|44623|{CHAVE_ENTRADA}"
                "|01052021|01052021|200,00|2|0|0|200,00|9|0|0|0|0|0|0|0|0|0|0|0|0|")
C170_1 = ("|C170|1|1000144|Iogurte|10|CX|100,00|0|0|010|1403|300|100,00|18|18,00"
          "|150,00|18|9,00||49||0|0|0||0|0|0|0|0||0|0|0|0|0|1.1.20.01.01|0|")
C170_2 = ("|C170|2|537861|Abobora|20,5|KG|100,00|0|0|060|1403|300|0|0|0"
          "|0|0|0||49||0|0|0||0|0|0|0|0||0|0|0|0|0|1.1.20.01.01|0|")
C170_SEM_CADASTRO = ("|C170|3|999999|Sem cadastro|1|UN|5,00|0|0|060|1403|300|0|0|0"
                     "|0|0|0||49||0|0|0||0|0|0|0|0||0|0|0|0|0|1.1.20.01.01|0|")
C190_ENTRADA = "|C190|010|1403|18|100,00|100,00|18,00|150,00|9,00|0|0||"
C190_ENTRADA_60 = "|C190|060|1403|0|105,00|0|0|0|0|0|0||"

# saída própria: sem C170, só o analítico — como nos arquivos reais
C100_SAIDA = (f"|C100|1|0|C001|55|00|001|44624|{CHAVE_SAIDA}"
              "|02052021|02052021|300,00|2|0|0|300,00|9|0|0|0|0|0|0|0|0|0|0|0|0|")
C190_SAIDA_60 = "|C190|060|5405|0|250,00|0|0|0|0|0|0||"
C190_SAIDA_00 = "|C190|000|5102|18|50,00|50,00|9,00|0|0|0|0||"

C800 = (f"|C800|59|00|62537|03052021|36,20|0|0||353552|{CHAVE_CUPOM}"
        "|0|36,20|0|0|0|0|")
C850 = "|C850|060|5405|0|36,2|0|0||"

H005 = "|H005|30042021|196,39|01|"
H010_A = "|H010|1000144|CX|81|1,19|96,39|0|||1.1.20.01.01|96,39|"
H010_B = "|H010|537861|KG|50|2,00|100,00|0|||1.1.20.01.01|100,00|"


def escrever(tmp_path, nome: str, linhas: list[str]) -> str:
    caminho = tmp_path / nome
    caminho.write_bytes(("\r\n".join(linhas) + "\r\n").encode("latin-1"))
    return str(caminho)


@pytest.fixture
def efds(tmp_path) -> list[str]:
    """Maio/2021 com tudo; abril/2021 só com o cadastro antigo e o inventário."""
    maio = escrever(tmp_path, "efd_2021_05.txt", [
        CABECALHO, ITEM_A_V2, CONV_A_FD, ITEM_B,
        C100_ENTRADA, C170_1, C170_2, C170_SEM_CADASTRO, C190_ENTRADA, C190_ENTRADA_60,
        C100_SAIDA, C190_SAIDA_60, C190_SAIDA_00,
        C800, C850,
    ])
    abril = escrever(tmp_path, "efd_2021_04.txt", [
        CABECALHO.replace("01052021|31052021", "01042021|30042021"),
        ITEM_A_V1, CONV_A_FD_ANTIGO, ITEM_B,
        H005, H010_A, H010_B,
    ])
    return [maio, abril]


@pytest.fixture
def conferidos(tmp_path) -> str:
    """A conferência achou o XML da entrada e do cupom; a saída ficou pendente."""
    caminho = str(tmp_path / "conferidos.parquet")
    pq.write_table(pa.table({"chave": [CHAVE_ENTRADA, CHAVE_CUPOM],
                             "tem_chave": [True, True]}), caminho)
    return caminho


@pytest.fixture
def extraido(efds, tmp_path):
    destino = str(tmp_path / "saida")
    progresso = extrair_movimentos(efds, destino)
    return progresso, destino


@pytest.fixture
def consolidado(extraido, conferidos):
    progresso, destino = extraido
    resumo = consolidar(destino, conferidos)
    return resumo, destino


def _linhas(caminho: str, sql: str = "SELECT * FROM r"):
    con = duckdb.connect()
    try:
        con.execute(f"CREATE VIEW r AS SELECT * FROM read_parquet('{caminho}')")
        return con.execute(sql).fetchall()
    finally:
        con.close()


class TestExtracao:
    def test_conta_o_que_leu(self, extraido):
        progresso, _ = extraido
        assert progresso.arquivos_lidos == 2
        assert progresso.documentos == 3            # 2 C100 + 1 C800
        assert progresso.documentos_com_item == 1   # só a entrada tem C170
        assert progresso.movimentos == 3
        assert progresso.analiticos == 5
        assert progresso.itens_cadastrados == 4     # 2 por período
        assert progresso.em_estoque == 2
        assert progresso.orfaos == 0
        assert progresso.recusados == []

    def test_o_item_sai_amarrado_ao_documento_pai(self, extraido):
        _, destino = extraido
        linhas = _linhas(f"{destino}/_movimentos_brutos.parquet",
                         "SELECT chave, operacao, data, codigo, quantidade, valor_st "
                         "FROM r ORDER BY numero_item")
        assert [l[0] for l in linhas] == [CHAVE_ENTRADA] * 3
        assert linhas[0][1:] == ("entrada", date(2021, 5, 1), "1000144",
                                 Decimal("10.00000"), Decimal("9.00"))
        assert linhas[1][3:5] == ("537861", Decimal("20.50000"))

    def test_documentos_dizem_quem_tem_item(self, extraido):
        _, destino = extraido
        linhas = dict((l[0], l[1:]) for l in _linhas(
            f"{destino}/{ARQUIVO_DOCUMENTOS}",
            "SELECT chave, operacao, itens, analiticos FROM r"))
        assert linhas[CHAVE_ENTRADA] == ("entrada", 3, 2)
        assert linhas[CHAVE_SAIDA] == ("saida", 0, 2)
        assert linhas[CHAVE_CUPOM] == ("saida", 0, 1)

    def test_inventario_leva_a_data_do_h005(self, extraido):
        _, destino = extraido
        linhas = _linhas(f"{destino}/{ARQUIVO_INVENTARIO}",
                         "SELECT codigo, data_inventario, quantidade, valor FROM r ORDER BY 1")
        assert linhas == [
            ("1000144", date(2021, 4, 30), Decimal("81.00000"), Decimal("96.39")),
            ("537861", date(2021, 4, 30), Decimal("50.00000"), Decimal("100.00")),
        ]

    def test_a_data_do_movimento_e_a_da_entrada_ou_saida_nao_a_emissao(self, tmp_path):
        """A Ficha 3 e o DATA do 1100 pedem a data da operação. A nota emitida em
        abril e escriturada em maio entra em maio; a saída própria sem DT_E_S
        fica com a emissão."""
        emitida_antes = C100_ENTRADA.replace("|01052021|01052021|", "|28042021|03052021|")
        saida_sem_dt_e_s = (C100_SAIDA.replace("|02052021|02052021|", "|02052021||")
                            .replace("|C100|1|0|C001|55|00|001|44624|", "|C100|1|0|C001|55|00|001|44625|"))
        efd = escrever(tmp_path, "efd.txt", [CABECALHO, ITEM_A_V2, emitida_antes, C170_1,
                                             saida_sem_dt_e_s, C170_1.replace("|1403|", "|5405|")])
        destino = str(tmp_path / "s")
        extrair_movimentos([efd], destino)
        linhas = _linhas(f"{destino}/_movimentos_brutos.parquet",
                         "SELECT operacao, data, serie FROM r ORDER BY operacao")
        assert linhas == [("entrada", date(2021, 5, 3), "001"), ("saida", date(2021, 5, 2), "001")]

    def test_item_antes_do_documento_e_orfao_e_fica_no_log(self, tmp_path):
        quebrado = escrever(tmp_path, "quebrado.txt", [CABECALHO, C170_1, C100_ENTRADA, C170_2])
        progresso = extrair_movimentos([quebrado], str(tmp_path / "s"))
        assert progresso.orfaos == 1
        assert progresso.movimentos == 1
        assert progresso.arquivos_com_orfaos == ["quebrado.txt"]

    def test_arquivo_ilegivel_nao_derruba_os_outros(self, efds, tmp_path):
        progresso = extrair_movimentos([str(tmp_path / "nao_existe.txt")] + efds,
                                       str(tmp_path / "s"))
        assert progresso.arquivos_lidos == 3
        assert len(progresso.recusados) == 1
        assert progresso.movimentos == 3


class TestConsolidacao:
    def test_vale_o_cadastro_mais_recente(self, consolidado):
        _, destino = consolidado
        linhas = dict(_linhas(f"{destino}/{ARQUIVO_ITENS}",
                              "SELECT codigo, descricao FROM r"))
        assert linhas == {"1000144": "Iog Vidativa 160g Ameixa",
                          "537861": "Abobora P/Doce Kg Frac"}

    def test_movimento_leva_cadastro_e_marca_da_conferencia(self, consolidado):
        _, destino = consolidado
        linhas = _linhas(
            f"{destino}/{ARQUIVO_MOVIMENTOS}",
            "SELECT codigo, descricao, ncm, cest, classificacao FROM r ORDER BY codigo")
        assert linhas == [
            ("1000144", "Iog Vidativa 160g Ameixa", "04031000", "1702200", "conferido"),
            ("537861", "Abobora P/Doce Kg Frac", "07099990", "", "conferido"),
            ("999999", None, None, None, "conferido"),
        ]

    def test_analitico_sabe_se_o_documento_tem_item(self, consolidado):
        _, destino = consolidado
        linhas = _linhas(f"{destino}/{ARQUIVO_ANALITICO}",
                         "SELECT chave, cst_icms, valor_operacao, tem_item FROM r ORDER BY 1, 2")
        por_chave = {}
        for chave, cst, valor, tem in linhas:
            por_chave.setdefault(chave, []).append((cst, valor, tem))
        assert por_chave[CHAVE_ENTRADA] == [("010", Decimal("100.00"), True),
                                            ("060", Decimal("105.00"), True)]
        assert por_chave[CHAVE_SAIDA] == [("000", Decimal("50.00"), False),
                                          ("060", Decimal("250.00"), False)]
        assert por_chave[CHAVE_CUPOM] == [("060", Decimal("36.20"), False)]

    def test_vale_o_fator_de_conversao_mais_recente_e_amarrado_ao_item(self, consolidado):
        """O 0220 não repete o código: é do 0200 logo acima. E, como o
        cadastro, vale o do período mais recente."""
        _, destino = consolidado
        linhas = _linhas(f"{destino}/conversoes.parquet", "SELECT codigo, unidade, fator FROM r")
        assert linhas == [("1000144", "FD", Decimal("12.000000000"))]

    def test_os_intermediarios_somem(self, consolidado):
        import os  # noqa: PLC0415

        _, destino = consolidado
        assert not any(n.startswith("_") for n in os.listdir(destino))


class TestResumo:
    def test_numeros(self, consolidado):
        r, _ = consolidado
        assert r.arquivos == 1                     # abril não tem documento
        assert r.documentos == 3
        assert r.documentos_com_item == 1
        assert r.entradas_sem_item == 0
        assert r.entradas_proprias_sem_item == 0
        assert r.saidas_sem_item == 2
        assert r.valor_saidas_sem_item == Decimal("336.20")      # 300 + 36,20
        assert r.valor_saidas_sem_item_st == Decimal("286.20")   # 250 + 36,20
        assert r.movimentos == 3
        assert r.movimentos_entrada == 3
        assert r.valor_entradas == Decimal("205.00")
        assert r.st_nas_entradas == Decimal("9.00")
        assert r.itens_cadastrados == 2
        assert r.itens_movimentados == 3
        assert r.itens_sem_cadastro == 1
        assert r.inventarios == 1
        assert r.itens_em_estoque == 2
        assert r.valor_em_estoque == Decimal("196.39")
        assert r.conferencia_usada
        assert r.cobertura_de_item == pytest.approx(1 / 3)

    def test_recortes(self, consolidado):
        r, _ = consolidado
        assert {f.codigo: f.documentos for f in r.por_cst} == {"010": 1, "060": 2}
        assert [(f.codigo, f.documentos) for f in r.por_classificacao] == [("conferido", 3)]
        assert [(f.codigo, f.documentos) for f in r.saidas_sem_item_por_modelo] == [
            ("55", 1), ("59", 1)]

    def test_avisos_dizem_o_que_a_efd_nao_tem(self, consolidado):
        r, _ = consolidado
        avisos = "\n".join(r.avisos)
        assert "2 documento(s) de saída (NF-e, CF-e-SAT) não trazem item" in avisos
        assert "R$ 336,20 de saídas" in avisos
        assert "R$ 286,20 com CST 60" in avisos
        assert "1 código(s) movimentados não estão no cadastro" in avisos
        assert "pendentes" not in avisos          # tudo o que tem item foi conferido

    def test_sem_conferencia_marca_e_avisa(self, extraido):
        _, destino = extraido
        r = consolidar(destino, None)
        assert not r.conferencia_usada
        assert [(f.codigo, f.documentos) for f in r.por_classificacao] == [
            ("nao_conferido", 3)]
        assert any("não foram marcados" in a for a in r.avisos)

    def test_movimento_de_documento_pendente_fica_marcado(self, extraido, tmp_path):
        _, destino = extraido
        vazio = str(tmp_path / "nenhum.parquet")
        pq.write_table(pa.table({"chave": pa.array([], pa.string())}), vazio)
        r = consolidar(destino, vazio)
        assert [(f.codigo, f.documentos) for f in r.por_classificacao] == [("pendente", 3)]
        assert any("3 movimento(s) são de documentos ainda pendentes" in a
                   for a in r.avisos)


class TestAvisosDoDominio:
    def test_entrada_propria_sem_item_e_regra_e_terceiros_e_anomalia(self):
        # base real: 179.333 entradas sem C170, todas de emissão própria
        r = ResumoDaMovimentacao(entradas_proprias_sem_item=179_333,
                                 entradas_sem_item=0, conferencia_usada=True)
        avisos = "\n".join(r.avisos)
        assert "179.333 nota(s) de entrada de emissão própria" in avisos
        assert "O item virá do XML" in avisos
        assert "de terceiros" not in avisos

        r = ResumoDaMovimentacao(entradas_sem_item=7, conferencia_usada=True)
        assert any("7 entrada(s) de terceiros não trazem C170" in a for a in r.avisos)

    def test_entrada_propria_sem_c170_conta_separada(self, tmp_path):
        # uma NF-e de entrada emitida pela própria empresa, sem item, ao lado
        # de uma entrada de terceiros com item
        propria = (f"|C100|0|0|C001|55|00|001|44700|{CHAVE_SAIDA}"
                   "|02052021|02052021|30,00|2|0|0|30,00|9|0|0|0|0|0|0|0|0|0|0|0|0|")
        efd = escrever(tmp_path, "efd.txt", [CABECALHO, ITEM_A_V2, C100_ENTRADA, C170_1,
                                             C190_ENTRADA, propria, C190_SAIDA_00])
        destino = str(tmp_path / "s")
        extrair_movimentos([efd], destino)
        r = consolidar(destino, None)
        assert r.documentos == 2
        assert r.entradas_proprias_sem_item == 1
        assert r.entradas_sem_item == 0
        assert r.saidas_sem_item == 0

    def test_sem_inventario(self):
        r = ResumoDaMovimentacao(movimentos=10, inventarios=0, conferencia_usada=True)
        assert any("Nenhum inventário" in a for a in r.avisos)

    def test_varios_estabelecimentos(self):
        r = ResumoDaMovimentacao(estabelecimentos=["1", "2", "3", "4"],
                                 conferencia_usada=True)
        aviso = next(a for a in r.avisos if "estabelecimentos" in a)
        assert "4 estabelecimentos" in aviso and "e mais 1" in aviso

    def test_saidas_sem_item_com_fatias(self):
        r = ResumoDaMovimentacao(
            saidas_sem_item=5, valor_saidas_sem_item=Decimal("1234.50"),
            valor_saidas_sem_item_st=Decimal("1000"),
            saidas_sem_item_por_modelo=[Fatia("NF-e", 5, codigo="55")],
            conferencia_usada=True)
        aviso = next(a for a in r.avisos if "não trazem item" in a)
        assert "R$ 1.234,50" in aviso and "R$ 1.000,00" in aviso


class TestPlanilhas:
    def test_as_quatro_saem(self, consolidado, tmp_path):
        _, destino = consolidado
        assert gerar_movimentos(f"{destino}/{ARQUIVO_MOVIMENTOS}",
                                str(tmp_path / "m.xlsx")) == 3
        assert gerar_itens(f"{destino}/{ARQUIVO_ITENS}", str(tmp_path / "i.xlsx")) == 2
        assert gerar_inventario(f"{destino}/{ARQUIVO_INVENTARIO}",
                                str(tmp_path / "e.xlsx")) == 2
        assert gerar_analitico(f"{destino}/{ARQUIVO_ANALITICO}",
                               str(tmp_path / "a.xlsx")) == 5

    def test_quantidade_com_fracao_e_booleano_em_portugues(self, consolidado, tmp_path):
        import openpyxl  # noqa: PLC0415

        _, destino = consolidado
        m = str(tmp_path / "m.xlsx")
        gerar_movimentos(f"{destino}/{ARQUIVO_MOVIMENTOS}", m)
        aba = openpyxl.load_workbook(m).active
        colunas = [c.value for c in aba[1]]
        quantidades = [aba.cell(row=i, column=colunas.index("Quantidade") + 1).value
                       for i in (2, 3, 4)]
        assert 20.5 in quantidades
        marca = aba.cell(row=2, column=colunas.index("Conferência") + 1).value
        assert marca == "Documento conferido"

        a = str(tmp_path / "a.xlsx")
        gerar_analitico(f"{destino}/{ARQUIVO_ANALITICO}", a)
        aba = openpyxl.load_workbook(a).active
        colunas = [c.value for c in aba[1]]
        valores = {aba.cell(row=i, column=colunas.index("Documento tem item na EFD") + 1).value
                   for i in range(2, 7)}
        assert valores == {"Sim", "Não"}

    def test_filtro_por_classificacao(self, extraido, tmp_path):
        _, destino = extraido
        consolidar(destino, None)
        assert gerar_movimentos(f"{destino}/{ARQUIVO_MOVIMENTOS}", str(tmp_path / "f.xlsx"),
                                classificacoes=frozenset({"conferido"})) == 0
        assert gerar_movimentos(f"{destino}/{ARQUIVO_MOVIMENTOS}", str(tmp_path / "g.xlsx"),
                                classificacoes=frozenset({"nao_conferido"})) == 3
